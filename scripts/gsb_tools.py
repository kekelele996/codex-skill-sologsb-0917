#!/usr/bin/env python3
"""GSB validation, official schema handling, and local Excel export."""
from __future__ import annotations

import difflib
import json
import os
import re
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

import device_config as _device_config
from common import (
    SCHEMA_FALLBACK,
    SOLO_SCRIPTS,
    SologsbError,
    atomic_write_text,
    commit_url,
    ensure_single_side_trace,
    read_json,
    save_state,
    sha256_text,
    text_non_whitespace_len,
    utc_now,
    write_json,
)
from evidence import run_audit

def gsb_server() -> str:
    """SOLO2 平台地址。技能包里不留默认域名，必须由设备配置提供。"""
    value = os.environ.get("SOLO2_SERVER", "").strip().rstrip("/")
    if not value:
        raise SologsbError(
            "缺少 SOLO2 平台地址：请在设备配置里设置 solo2.baseUrl，"
            "或运行 scripts/configure.py wizard"
        )
    return value
AUDIT_HUMAN = SOLO_SCRIPTS / "audit-human-writing.py"
EXPECTED_FINGERPRINT = "954e9db2d25afeb4"
FILE_TOKEN = re.compile(r"[A-Za-z0-9_./-]+\.(?:go|js|cjs|mjs|ts|tsx|jsx|py|java|kt|rs|vue|json|ya?ml|toml|md|sql|sh|css|html|xml)")
ERROR_TOKEN = re.compile(r"(?:Error|ERROR|panic|PANIC|npm ERR!|failed|FAILED|报错|失败)[:：]?\s*[^\n，。；;]{1,120}")
TEST_COUNT_PATTERN = re.compile(
    r"(?:(?:后端|前端|API|接口|测试|断言|用例)[^。；，]{0,12}?\d+\s*(?:个|项|条|步|轮)|"
    r"\d+\s*(?:个|项|条|步|轮)[^。；，]{0,12}?(?:后端|前端|API|接口|测试|断言|用例))",
    re.I,
)
MARKDOWN_REASON_PATTERNS = (
    ("代码块或行内代码", re.compile(r"```|`")),
    ("分隔线", re.compile(r"(?m)^\s{0,3}(?:-{3,}|\*{3,}|_{3,})\s*$")),
    ("链接或图片", re.compile(r"!?\[[^\]]+\]\([^)]+\)")),
    ("强调标记", re.compile(r"(?:\*\*|__|~~|(?<!\*)\*[^*\n]+\*(?!\*)|(?<!_)_[^_\n]+_(?!_))")),
    ("HTML 标签", re.compile(r"</?[A-Za-z][^>]*>")),
    (
        "标题、引用、列表或表格",
        re.compile(r"(?m)^\s{0,3}(?:#{1,6}\s|>\s?|[-+*]\s|\d+[.)]\s|\|.*\|\s*$)"),
    ),
)
LOW_VALUE_PROCESS_NOISE_PATTERN = re.compile(
    r"(?:String to replace not found|文本替换失败|替换文本未找到|编辑未匹配|"
    r"No module named|ModuleNotFoundError|command not found|命令未找到|"
    r"(?:python|python3|node|npm|pnpm|git|bash)\s*(?:命令)?\s*(?:未找到|缺失|返回\s*127)|"
    r"退出码\s*127)",
    re.I,
)
FIELD_FACTOR_RE = re.compile(
    r"(?:录屏|屏幕录制|录制过程|录制画面|录制结果|视频画面|视频中|视频里|视频显示|视频可见|视频证据|"
    r"视频|截图|画面中|画面显示|画面可见|镜头|剪辑|剪掉|\bOtty\b|\biTerm2?\b|\bTerminal\.app\b|"
    r"1280\s*[x×]\s*720|720p|\bMP4\b|鼠标|光标|终端窗口|终端界面|命令行窗口|"
    r"浏览器|浏览器窗口|屏幕|测试设备|测试机|运行环境|运行机器|验收宿主|验收机|采集环境|采集设备|录制设备)",
    re.I,
)
NON_CONTAINER_TEST_ARTIFACT_RE = re.compile(
    r"(?:本地|容器外|非容器内|宿主(?:机|侧)?)[^。；，]{0,16}"
    r"(?:测试文件|测试脚本|测试代码|自测文件|自测脚本|验收脚本|验收测试文件|验收测试脚本)|"
    r"(?:验收脚本|验收测试文件|验收测试脚本)|"
    r"(?:测试文件|测试脚本|测试代码|自测文件|自测脚本)[^。；，]{0,16}"
    r"(?:由本地|在容器外|非容器内|不是容器内|容器外生成|本地生成)",
    re.I,
)
BANNED_REASON_PATTERN = re.compile(r"落在.{0,16}")
# 2026-09-21 起：文案通则里的禁用词与排版规则统一从 references 读取。
# SKILL_ROOT: 本文件位于 <skill>/scripts/gsb_tools.py
SKILL_ROOT = Path(__file__).resolve().parents[1]
REASON_WORD_REPLACEMENTS = SKILL_ROOT / "references" / "reason-word-replacements.json"
REASON_LABELS = ("A 侧方案", "B 侧方案")
# 数字、英文与中文之间不留空格：写“计数从2变22”“APIs.vue的回调”。
REASON_SPACING_PATTERNS = (
    re.compile(r"([A-Za-z0-9]+)(\s+)(?=[\u4e00-\u9fff])"),
    re.compile(r"(?<=[\u4e00-\u9fff])(\s+)([A-Za-z0-9]+)"),
)
REASON_SINGLE_LETTER_LABEL = re.compile(r"^[A-Za-z]$")
# 文风“朴实但不干练”：语气补足词与句均长度。
REASON_SOFTENER_PATTERN = re.compile(r"(还是|仍然|已经|随后|后来|之后|了|过|着|掉|各加一|多搬|就|都|才|又|再)")
REASON_MIN_AVG_SENTENCE = 13
REASON_MIN_SOFTENERS = 3
REASON_MAX_SENTENCE_CHARS = 56
REASON_MAX_SENTENCE_COMMAS = 5
REASON_VAGUE_PATTERN = re.compile(r"(?:真实|真正|其实|本质上|实际上)")
# 2026-09-23 起：平台“理由 AI 化打分”会扣句句同主语起头、碎句和“只罗列不交代判准”。
REASON_MAX_LABEL_MENTIONS = 3
REASON_MIN_SENTENCE_CHARS = 8
REASON_CRITERION_PATTERN = re.compile(
    r"(?:看重|要紧|关键|决定|差别在|差距在|分开|分出|拉开|主要看|更在意|优先|首先要|最重要)"
)
# 2026-09-24 起：平台以“电报体”打回（八句一句一事实、句号密集、因果和扣分点靠读者自己排）。
REASON_MAX_SENTENCES = 6
REASON_MIN_LINK_WORDS = 2
REASON_LINK_PATTERN = re.compile(
    r"(?:因为|由于|所以|因而|于是|结果|导致|以致|使得|这样一来|这就|但是|但|不过|可是|却|而是|只要|一旦|否则|才)"
)
# 2026-09-25 起：为了避开“相邻句同称谓起头”，理由出现“检查……后，……”这类无主语句，
# 读者分不清是哪一侧做的。打平统一写“Same”，但“选A”“选B”这类选项代号不进正文。
REASON_SUBJECTLESS_OPENER = re.compile(
    r"^(?:检查|核对|读取|阅读|查看|查阅|梳理|比对|运行|执行|调用|修改|编辑|调整|排查|定位|验证|构建|启动|安装|补充|补齐|新增|删除|实现)"
)
REASON_FORM_OPTION_PATTERN = re.compile(r"选\s*[AB](?![A-Za-z\s]*侧)")
# “并发保存先提交标题未更新”一类：省掉主语和衔接，把两件事压成一个短语，要回读才懂。
REASON_COMPRESSED_CLAUSE = re.compile(r"先[\u4e00-\u9fffA-Za-z0-9]{1,8}未[\u4e00-\u9fff]{1,6}")
REASON_FLUENCY_PATTERNS = (
    ("重复标点", re.compile(r"[，。；：！？!?]{2,}")),
    ("重复虚词", re.compile(r"(?:的的|了了|是是|在在|和和|与与|就就|都都)")),
    (
        "连接词堆叠",
        re.compile(
            r"(?:但是|然而|并且|而且|同时|另外|此外).{0,18}"
            r"(?:但是|然而|并且|而且|同时|另外|此外)"
        ),
    ),
    (
        # “都成功了。”“出错了。”是口语复盘最自然的收尾，不算残句（2026-09-25）。
        "残句结尾",
        re.compile(
            r"(?:的|和|与|及|而|但|因为|由于|如果|当|在|从|对|把|被|为|是|就|都|还|也|很|更|最|可以|能够|需要|应该|必须)[。！？!?]$"
        ),
    ),
)
EVALUATION_EXCLUDED_NOISE_RE = re.compile(
    r"(?:String to replace not found|文本替换失败|替换文本未找到|编辑未匹配|"
    r"No module named|ModuleNotFoundError|command not found|命令未找到|"
    r"(?:python|python3|node|npm|pnpm|git|bash)\s*(?:命令)?\s*(?:未找到|缺失|返回\s*127)|"
    r"退出码\s*127|临时工作目录|工作目录不存在|测试\s*PYTHONPATH\s*缺失)",
    re.I,
)


def banned_reason_words() -> dict[str, str]:
    """Return the shared banned-word map (single source: references/reason-word-replacements.json)."""
    fallback = {"落库": "入库", "闭环": "完整覆盖", "根因": "原因"}
    try:
        data = json.loads(REASON_WORD_REPLACEMENTS.read_text(encoding="utf-8"))
    except Exception:
        return fallback
    words = data.get("replace") if isinstance(data, dict) else None
    if not isinstance(words, dict) or not words:
        return fallback
    return {str(key): str(value) for key, value in words.items()}


def reason_style_errors(reason: str) -> list[str]:
    """Blocking style rules from references/reason-writing-rules.md."""
    errors: list[str] = []
    for word, better in banned_reason_words().items():
        if word and word in reason:
            errors.append(f"GSB 理由禁用“{word}”，请改写成“{better}”: {word}")
    for pattern in REASON_SPACING_PATTERNS:
        for match in pattern.finditer(reason):
            token = match.group(1) if match.group(1).strip() else match.group(2)
            if REASON_SINGLE_LETTER_LABEL.match(token):
                # “A 侧方案”“B 侧方案”是固定称谓，允许保留一个空格。
                continue
            errors.append(
                "GSB 理由的数字与英文两侧不要留空格，直接连写，例如“计数从2变22”“APIs.vue的回调”: "
                f"{match.group(0)!r}"
            )
            break
        else:
            continue
        break
    return errors


def reason_style_warnings(reason: str) -> list[str]:
    """Non-blocking wording hints: plain language should not read like telegraphic notes."""
    warnings: list[str] = []
    sentences = [part for part in re.split(r"[。；！？!?]", reason) if part.strip()]
    if sentences:
        average = sum(len(part) for part in sentences) / len(sentences)
        if average < REASON_MIN_AVG_SENTENCE:
            warnings.append(
                "GSB 理由读起来偏干练，句子过短；请把动词后的语气和结果补足，"
                f"例如“建了独立表”“覆盖掉了”: 句均 {average:.1f} 字"
            )
    softeners = len(REASON_SOFTENER_PATTERN.findall(reason))
    if softeners < REASON_MIN_SOFTENERS:
        warnings.append(
            "GSB 理由缺少语气与衔接词，读起来像提纲；建议补“还是照旧”“已经”“从…变…”这类说法: "
            f"命中 {softeners} 处"
        )
    vague = REASON_VAGUE_PATTERN.findall(reason)
    if vague:
        warnings.append(
            "GSB 理由出现“真实”“真正”“其实”等空泛表达；请改成能核对的动作、状态或结果: "
            + "、".join(dict.fromkeys(vague))
        )
    return warnings


def reason_flow_errors(reason: str) -> list[str]:
    """Block list-like reasons: 平台 AI 化打分会把“句句以 A 侧方案起头”的电报体判为 AI 文风。"""
    errors: list[str] = []
    sentences = [part.strip() for part in re.findall(r"[^。！？!?]+[。！？!?]?", reason) if part.strip()]
    previous = ""
    for index, sentence in enumerate(sentences, 1):
        label = next((item for item in REASON_LABELS if sentence.startswith(item)), "")
        if label and label == previous:
            errors.append(
                f"GSB 理由第 {index - 1}、{index} 句都以“{label}”起头，读起来像逐条清单；"
                "后一句直接接着说，或用“随后”“这一改动”这类承接"
            )
            break
        previous = label
    for index, sentence in enumerate(sentences, 1):
        opener = REASON_SUBJECTLESS_OPENER.match(sentence)
        if opener:
            errors.append(
                f"GSB 理由第 {index} 句以“{opener.group(0)}”起头，没有交代是哪一侧做的，"
                "读者会把它算到上一侧；换侧时句首写明称谓，同一侧接着说用“随后”“它”承接"
            )
            break
    form_option = REASON_FORM_OPTION_PATTERN.search(reason)
    if form_option:
        errors.append(
            f"GSB 理由把表单选项“{form_option.group(0)}”写进了正文；结论用中文说，"
            "例如“因此选择Same”“因此选择 B 侧方案”"
        )
    for label in REASON_LABELS:
        count = reason.count(label)
        if count > REASON_MAX_LABEL_MENTIONS:
            errors.append(
                f"“{label}”出现 {count} 次，超过 {REASON_MAX_LABEL_MENTIONS} 次；"
                "同一侧的事实连成一段写，不要每句重复主语"
            )
    for index, sentence in enumerate(_reason_sentences(reason), 1):
        clean = re.sub(r"\s+", "", sentence)
        if len(clean) < REASON_MIN_SENTENCE_CHARS:
            errors.append(
                f"GSB 理由第 {index} 句只有 {len(clean)} 字，像补在末尾的碎句：{clean}；并入前后句"
            )
    if len(sentences) > REASON_MAX_SENTENCES:
        errors.append(
            f"GSB 理由共 {len(sentences)} 句，超过 {REASON_MAX_SENTENCES} 句，句号过密像电报体；"
            "同一侧的动作和结果用逗号连成一句，交代清谁因谁果"
        )
    body = "".join(sentences[:-1]) if len(sentences) > 1 else reason
    links = REASON_LINK_PATTERN.findall(body)
    if len(links) < REASON_MIN_LINK_WORDS:
        errors.append(
            f"GSB 理由结论句之前只有 {len(links)} 处因果或转折衔接，读者要自己排谁因谁果；"
            "负面事实用“结果”“导致”“所以”接上后果，两侧对照用“但”“不过”“却”"
        )
    compressed = REASON_COMPRESSED_CLAUSE.search(reason)
    if compressed:
        errors.append(
            f"GSB 理由把两件事压成了一个短语：{compressed.group(0)}；补上主语和衔接，"
            "例如“两个人同时保存时，先提交的一方写进去了，标题却没有更新”"
        )
    if text_non_whitespace_len(reason) >= 150 and not REASON_CRITERION_PATTERN.search(reason):
        errors.append(
            "GSB 理由只罗列事实，没有交代哪一条是扣分点、最看重什么；结论前用一句话点明，"
            "例如“这个任务最重要的是并发保存不丢数据”"
        )
    return errors


def _reason_sentences(reason: str) -> list[str]:
    return [part.strip() for part in re.findall(r"[^。；！？!?]+[。；！？!?]?", reason) if part.strip()]


def _punctuation_pair_errors(sentence: str) -> list[str]:
    errors: list[str] = []
    for left, right, label in (
        ("（", "）", "圆括号"),
        ("(", ")", "半角圆括号"),
        ("《", "》", "书名号"),
        ("“", "”", "引号"),
    ):
        if sentence.count(left) != sentence.count(right):
            errors.append(f"{label}没有成对出现")
    return errors


def reason_language_errors(reason: str) -> list[str]:
    """Block overlong, incomplete or visibly disfluent Chinese sentences."""
    errors: list[str] = []
    stripped = reason.strip()
    if stripped and stripped[-1] not in "。！？!?":
        errors.append("GSB 理由最后一句缺少句末标点，句子没有收完整")
    for index, sentence in enumerate(_reason_sentences(reason), 1):
        clean = re.sub(r"\s+", "", sentence)
        if len(clean) > REASON_MAX_SENTENCE_CHARS:
            errors.append(
                f"GSB 理由第 {index} 句过长，共 {len(clean)} 字；单句不得超过 {REASON_MAX_SENTENCE_CHARS} 字"
            )
        comma_count = clean.count("，") + clean.count(",")
        if comma_count > REASON_MAX_SENTENCE_COMMAS:
            errors.append(
                f"GSB 理由第 {index} 句分句过多，共 {comma_count + 1} 个分句；请拆成通顺短句"
            )
        pair_errors = _punctuation_pair_errors(clean)
        if pair_errors:
            errors.append(f"GSB 理由第 {index} 句标点不通顺：{'；'.join(pair_errors)}")
        for label, pattern in REASON_FLUENCY_PATTERNS:
            match = pattern.search(clean)
            if match:
                errors.append(
                    f"GSB 理由第 {index} 句存在{label}：{match.group(0)[:30]}；请改成高中语文水平能直接读懂的通顺句子"
                )
    return errors


def _evaluation_excluded(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    value = item.get("evaluationExcluded")
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "excluded"}
    if isinstance(value, list):
        return bool(value)
    if isinstance(value, dict):
        return any(bool(value.get(key)) for key in ("excluded", "isExcluded", "value", "status"))
    return False


def evaluation_excluded_reason_errors(
    reason: str,
    draft: dict[str, Any],
    evidence_doc: dict[str, Any],
) -> list[str]:
    """Reject any reference to evidence that is excluded from evaluation."""
    excluded = [
        item for item in evidence_doc.get("evidence") or []
        if _evaluation_excluded(item)
    ]
    if not excluded:
        return []
    excluded_ids = {str(item.get("id") or "") for item in excluded}
    errors: list[str] = []
    for index, claim in enumerate(draft.get("claims") or [], 1):
        if not isinstance(claim, dict):
            continue
        for evidence_id in claim.get("evidenceIds") or []:
            if str(evidence_id) in excluded_ids:
                errors.append(
                    f"claim {index} 不得引用 evaluationExcluded 的环境或工具噪声证据 {evidence_id}"
                )
    for index, mapping in enumerate(draft.get("sentenceEvidence") or [], 1):
        if not isinstance(mapping, dict):
            continue
        for evidence_id in mapping.get("evidenceIds") or []:
            if str(evidence_id) in excluded_ids:
                errors.append(
                    f"sentenceEvidence {index} 不得引用 evaluationExcluded 的环境或工具噪声证据 {evidence_id}"
                )
    clean_reason = _normalize(reason)
    for item in excluded:
        text = _normalize(str(item.get("text") or ""))
        if len(text) >= 6 and text in clean_reason:
            errors.append(
                "GSB 理由不得写入 evaluationExcluded 的环境或工具噪声事实："
                f"{str(item.get('text') or '')[:60]}"
            )
    return list(dict.fromkeys(errors))


TRIGGER_KINDS = {"step", "file", "command", "requirement"}
OBJECTIVE_CONSEQUENCE_RE = re.compile(
    r"(导致|造成|没有|未|失败|报错|退出码|无法|阻断|拒绝|缺少|缺失|未落地|未启动|未接入|未完成|影响)"
)
MACHINE_STEP_RE = re.compile(r"第\s*\d{3,}\s*次")
GENERIC_TRIGGER_RE = re.compile(r"^(?:过程(?:中|里)?|操作(?:时|中)|某一步|执行(?:时|中)|运行(?:时|中))$")
STEP_TRIGGER_RE = re.compile(
    r"(第\s*[一二三四五六七八九十\d]+\s*(?:步|次|轮|阶段)|步骤\s*\d+|初始化(?:阶段|时)?|"
    r"首次(?:启动|运行|执行|提交|联调|验证)(?:阶段|时)?|"
    r"第一次(?:启动|运行|执行|提交|联调|验证)(?:阶段|时)?|启动阶段|运行阶段|"
    r"联调(?:阶段|时)?|验证阶段|复跑(?:阶段|时)?|返修(?:阶段|时)?|重跑(?:阶段|时)?)"
)
NUMBERED_STEP_RE = re.compile(r"第\s*[一二三四五六七八九十\d]+\s*(?:步|次|轮|阶段)")
STEP_ACTION_RE = re.compile(
    r"执行|调用|运行|修改|新增|编辑|提交|启动|测试|构建|修复|重命名|导入|导出|联调|验证|复跑"
)
COMMAND_TRIGGER_RE = re.compile(
    r"`[^`]+`|\b(?:go|npm|pnpm|yarn|pytest|python|node|docker|mvn|gradle|git|curl|make|bash)\b"
)
REQUIREMENT_TRIGGER_RE = re.compile(
    r"(?:(?:新增|编辑|重命名|删除|查询|导入|导出|修改|同步).{0,20}"
    r"(?:时|阶段|场景|环节|需求|要求)|(?:需求|用例|提示词|验收).{0,30})"
)

PROCESS_ACTION_RE = re.compile(
    r"(?:阅读|读取|查阅|浏览|查看|检查|核对|定位|搜索|检索|梳理|比对|修改|编辑|调整|重构|"
    r"补充|补齐|补全|同步|运行|执行|调用|调试|排查|验证|测试|构建|启动|安装|新增|删除|实现|发现|修复)"
)
PROCESS_LOCATOR_RE = re.compile(
    r"(?:第\s*[一二三四五六七八九十\d]+\s*(?:步|次|轮|阶段)|步骤\s*\d+|阶段|"
    r"初始化|联调|验证|构建|启动|迁移|数据库|脚本|文件|接口|服务|模块|路由|中间件|指标|功能|字段|"
    r"页面|流程|逻辑|配置|模型|需求|用例|`[^`]+`|\.(?:go|js|cjs|mjs|ts|tsx|jsx|py|java|kt|rs|vue|json|ya?ml|toml|md|sql|sh|css|html|xml))"
)
ARTIFACT_OUTCOME_RE = re.compile(
    r"(?:返回|输出|缺少|缺失|未实现|没有|失败|报错|异常|保留|仍然|仍含|写入|生成|创建|"
    r"删除|更新|展示|完成|支持|正常|通过|可用|实现|拒绝|500|404)"
)

# 交付完整性描述（a/b_desc_delivery）的门禁口径，详见 references/delivery-scoring.md。
DELIVERY_DESC_MIN = 40
DELIVERY_DESC_MAX = 200
DELIVERY_REASON_COPY_FRAGMENT = 20
DELIVERY_REASON_COPY_RATIO = 0.6
DELIVERY_SIDE_FRAGMENT = 12
DELIVERY_SIDE_RATIO = 0.5
DELIVERY_PROCESS_DIMENSION_RE = re.compile(
    r"(?:任务规划|规划能力|推理|思考过程|工具调用|ToolCall|TodoWrite|指令遵循|边界感|执行能力)", re.I
)
DELIVERY_VERIFY_RE = re.compile(r"(?:核对|验证|复跑|跑通|构建|启动|回读|通过|正常|一致)")
DELIVERY_UNRESOLVED_RE = re.compile(r"(?:未实现|没有实现|无法运行|无法启动|报错|缺少|缺失|遗漏|虚假成功)")
# 口语化的客观后果：“进不了”“没法验证”“返回404”也算写出了后果。
DELIVERY_CONSEQUENCE_RE = re.compile(r"(?:没法|进不了|打不开|用不了|跑不起来|不能|返回\s*[45]\d\d|读回为空|丢失)")
# 描述与轨迹一致性（红线）。
DELIVERY_RUNNABILITY_RE = re.compile(
    r"(?:build|构建|编译|compile|tsc|start|启动|serve|dev|up|install|安装|启动验证|探活|录制)", re.I
)
DELIVERY_RUN_OK_RE = re.compile(r"(?:构建(?:和启动)?(?:都)?通过|启动(?:都)?(?:正常|成功|通过)|一次性?跑通|均正常|都能走通|全部通过)")
DELIVERY_RUN_FAIL_RE = re.compile(r"(?:无法运行|无法启动|启动失败|构建失败|编译失败|跑不起来|起不来)")
# Python 的 \b 把中文也当单词字符，锚点紧贴中文时要用 ASCII 边界。
DELIVERY_LITERAL_ANCHOR_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?:[A-Z][A-Z0-9]*_[A-Z0-9_]{2,}|SQLSTATE\s*\d{5}|[A-Za-z]+(?:Error|Exception)|"
    r"(?:npm|pnpm|yarn|go|pytest|mvn|gradle|docker|cargo|make)\s+[a-z][A-Za-z0-9_:-]*)(?![A-Za-z0-9_])"
)
DELIVERY_OBSERVED_ANCHOR_RE = re.compile(r"(?:(?<![\d.])[45]\d\d(?![\d.])|/api(?:/[A-Za-z0-9_{}:.-]+)+)")
DELIVERY_COMPLETION_CLAIM_RE = re.compile(
    r"(?:完成|已实现|已修复|全部通过|均已|都已|done|implemented|all tests pass|successfully)", re.I
)
# 本地编写的测试与自动化脚本不参与交付完整性描述（2026-09-23 红线）。
DELIVERY_LOCAL_AUTOMATION_RE = re.compile(
    r"(?:自动化脚本|自动化测试|验收脚本|自测脚本|冒烟脚本|录制脚本|场景脚本|本地脚本|本地测试|"
    r"playwright|puppeteer|selenium|scenario|apiRequests)",
    re.I,
)
# 描述与理由都用主观直述（“请求了登录接口，返回404”），不交代信息从哪来。
SOURCE_ATTRIBUTION_RE = re.compile(
    r"(?:从(?:录屏|视频|截图|画面|测试|复核|验证|日志|结果)[^。；，]{0,8}(?:来看|看|可见|可知|得知)|"
    r"(?:录屏|视频|截图|画面)(?:里|中|上)?(?:显示|可见|看到|能看到|可以看到)|"
    r"(?:编写|写好|补写|新增|准备)的(?:测试|用例|脚本)|"
    r"(?:测试|复核|验证|探活|自测)(?:结果)?(?:显示|表明|证明|说明|可见)|"
    r"根据(?:测试|复核|验证|录屏|日志)|验证计划|探活|probe)",
    re.I,
)
# 轨迹里的改动记录：描述说“没改/未修改 X”却有编辑，或说“修改了 X”却从没碰过，都是完全对立。
DELIVERY_NOT_CHANGED_RE = re.compile(r"(?:没有|没|未|并未)(?:修改|改动|改过|动过|改)")
DELIVERY_CHANGED_RE = re.compile(r"(?:修改了|改了|改动了|重写了|新增了|补上了|加上了|接入了)")
EDIT_TOOL_NAMES = {"Edit", "Write", "MultiEdit", "NotebookEdit"}
DELIVERY_CLAIMED_RE = re.compile(r"(?:宣称|声称|称已|自称|回复说|总结里说|最终回复)")
DELIVERY_ACTUAL_RE = re.compile(r"(?:实际|diff|提交里|代码里|改动里|并未|没有改)", re.I)


def _keychain_secret(service: str) -> str:
    account = os.environ.get("USER", "")
    proc = subprocess.run(
        ["security", "find-generic-password", "-a", account, "-s", service, "-w"],
        capture_output=True,
        text=True,
        check=False,
    )
    return proc.stdout.strip() if proc.returncode == 0 else ""


def _request_json(path: str, *, timeout: int = 60) -> dict[str, Any]:
    def _call(cookie: str, csrf: str) -> dict[str, Any]:
        request = urllib.request.Request(
            gsb_server() + path,
            headers={
                "Accept": "application/json",
                "Cookie": cookie,
                "x-csrf-token": csrf,
                "Referer": gsb_server() + "/app/gsb/submit",
                "User-Agent": "sologsb-0917/1.0",
            },
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))

    service = os.environ.get("SOLOSB_SOLO2_KEYCHAIN_SERVICE", "").strip()
    cookie = os.environ.get("SOLO_QA_COOKIE", "").strip() or (
        _keychain_secret(service + "-cookie") if service else ""
    )
    csrf = os.environ.get("SOLO_QA_CSRF", "").strip() or (
        _keychain_secret(service + "-csrf") if service else ""
    )
    if not cookie or not csrf:
        # 配置里有账号密码时自动登录一次
        if _device_config.refresh_solo2_into_env():
            cookie = os.environ.get("SOLO_QA_COOKIE", "").strip()
            csrf = os.environ.get("SOLO_QA_CSRF", "").strip()
    if not cookie or not csrf:
        raise SologsbError("缺少 SOLO2 只读凭据")
    try:
        return _call(cookie, csrf)
    except urllib.error.HTTPError as exc:
        # 会话过期：用配置里的账号密码重登一次再重试
        if exc.code == 401 and _device_config.refresh_solo2_into_env():
            try:
                return _call(os.environ["SOLO_QA_COOKIE"], os.environ["SOLO_QA_CSRF"])
            except urllib.error.HTTPError as retry_exc:
                body = retry_exc.read().decode("utf-8", errors="replace")
                raise SologsbError(
                    f"重新登录后仍返回 HTTP {retry_exc.code}: {body[:1000]}"
                ) from retry_exc
        body = exc.read().decode("utf-8", errors="replace")
        raise SologsbError(f"只读请求 HTTP {exc.code}: {body[:1000]}") from exc


def load_official_schema(*, accept_change: bool = False) -> dict[str, Any]:
    fallback = read_json(SCHEMA_FALLBACK, {})
    try:
        current = _request_json("/api/v1/gsb/form-schema")
    except SologsbError as exc:
        if os.environ.get("SOLOSB_ALLOW_SCHEMA_FALLBACK") == "1":
            current = fallback
        else:
            raise SologsbError(f"无法读取官方 GSB schema，拒绝使用旧快照: {exc}") from exc
    if not current.get("fields"):
        raise SologsbError("官方 GSB schema 不可用")
    fingerprint = str(current.get("fingerprint") or "")
    if fingerprint != EXPECTED_FINGERPRINT and not accept_change:
        raise SologsbError(
            f"官方 GSB schema fingerprint 已变化: {fingerprint} != {EXPECTED_FINGERPRINT}；"
            "请检查字段后使用 --accept-schema-change"
        )
    return current


def _normalize(value: str) -> str:
    return re.sub(r"\s+", "", value).strip()


def _validate_artifact_description(reason: str) -> list[str]:
    """Keep GSB comparison focused on delivered product facts and natural wording."""
    errors: list[str] = []
    attribution = SOURCE_ATTRIBUTION_RE.search(reason)
    if attribution:
        errors.append(
            "GSB 理由不要交代信息来源，直接写做了什么、看到什么，例如“请求了登录接口，返回404”: "
            f"{attribution.group(0)}"
        )
    match = FIELD_FACTOR_RE.search(reason)
    if match:
        errors.append(
            "GSB 理由不得引用录屏、视频、截图、浏览器、测试设备、运行环境、验收宿主或采集现场因素；没有例外: "
            f"{match.group(0)}"
        )
    acceptance = NON_CONTAINER_TEST_ARTIFACT_RE.search(reason)
    if acceptance:
        errors.append(
            "GSB 理由不得引用非容器内生成的测试文件、测试脚本、自测脚本或验收脚本；"
            "这些本地（容器外）内容不参与 GSB 文案，没有例外: "
            f"{acceptance.group(0)}"
        )
    banned = BANNED_REASON_PATTERN.search(reason)
    if banned:
        errors.append(
            "GSB 理由禁止使用“落在……”式收束句式；请直接写“因此选择 B 侧方案”或“B 侧方案更好”: "
            f"{banned.group(0)}"
        )
    errors.extend(reason_style_errors(reason))
    return errors


SUBMISSION_SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "submission" / "scripts"


def _load_submission_preflight():
    # 历史查重与提交前 preflight 共用同一份缓存和判定口径；preflight 也会反向懒加载本模块，所以只能在函数里导入。
    import importlib

    if str(SUBMISSION_SCRIPTS_DIR) not in sys.path:
        sys.path.insert(0, str(SUBMISSION_SCRIPTS_DIR))
    return importlib.import_module("preflight")


def fetch_submission_history(task_root: Path | None = None) -> dict[str, Any]:
    """读取历史 GSB（本机共享缓存，过期才增量刷新），排除本任务自己的提交。"""
    preflight = _load_submission_preflight()
    current_sessions: set[str] = set()
    exclude_ids: set[int] = set()
    if task_root is not None:
        state = read_json(task_root / "monitor" / "state.json", {}) or {}
        current_sessions = {
            str(((state.get("sides") or {}).get(side) or {}).get("sessionId") or "")
            for side in ("A", "B")
        } - {""}
        result = read_json(task_root / "workspace" / "评审文件" / "pre-submit" / "submission-result.json", {}) or {}
        number = str(result.get("submissionNo") or "")
        if number.isdigit():
            exclude_ids = {int(number)}
    return preflight.fetch_gsb_prompt_history(current_sessions, exclude_ids)


def fetch_reason_history(task_root: Path | None = None) -> dict[str, Any]:
    history = fetch_submission_history(task_root)
    return {
        **history,
        "items": [item for item in history.get("items") or [] if str(item.get("gsbReason") or "").strip()],
    }


def _reason_similarity_approval_ok(task_root: Path, reason: str) -> bool:
    approval_path = task_root / "workspace" / "评审文件" / "pre-submit" / "reason-similarity-approval.json"
    if not approval_path.is_file():
        return False
    try:
        approval = json.loads(approval_path.read_text(encoding="utf-8"))
    except Exception:
        return False
    return (
        isinstance(approval, dict)
        and approval.get("approvalKind") == "manual"
        and approval.get("scope") == "reason-similarity"
        and str(approval.get("approvedBy") or "").strip().casefold()
        == (os.environ.get("SOLOGBS_AUTO_APPROVER", "").strip() or "auto").casefold()
        and approval.get("reviewDecision") == "REVIEW_REQUIRED"
        and str(approval.get("reasonSha256") or "") == sha256_text(reason.strip())
    )


def _validate_reason_dedup(reason: str, task_root: Path | None = None) -> list[str]:
    """G10：与提交前 preflight 的 B-5 同一口径，在草稿阶段就挡住历史理由复用。"""
    if task_root is not None and _reason_similarity_approval_ok(task_root, reason):
        return []
    try:
        history = fetch_reason_history(task_root)
        review = _load_submission_preflight().assess_gsb_reason_dedup(reason, history)
    except Exception as exc:
        if os.environ.get("SOLOSB_ALLOW_OFFLINE_DEDUP") == "1":
            return []
        return [f"无法读取历史 GSB 理由用于 G10 查重: {exc}"]
    decision = str(review.get("decision") or "")
    if decision in {"UNIQUE", "MISSING"}:
        return []
    top = (review.get("matches") or [{}])[0]
    return [
        f"G10 GSB 理由与历史理由重复（{decision}，最长公共片段 {top.get('longestCommonSubstringLength', 0)} 字）："
        f"{review.get('rewriteInstruction') or ''}"
    ]


def _validate_delivery_dedup(draft: dict[str, Any], task_root: Path | None = None) -> list[str]:
    """G12：交付完整性描述与历史 A/B 描述比对；只有 EXACT/SIMILAR 阻断，与 preflight 一致。"""
    delivery = draft.get("delivery") if isinstance(draft.get("delivery"), dict) else {}
    texts = {
        side: str(((delivery.get(side) or {}) if isinstance(delivery.get(side), dict) else {}).get("description") or "").strip()
        for side in ("A", "B")
    }
    if not all(texts.values()):
        return []  # 缺描述由 validate_delivery 报
    try:
        history = fetch_submission_history(task_root)
        review = _load_submission_preflight().assess_delivery_dedup(texts, history)
    except Exception as exc:
        if os.environ.get("SOLOSB_ALLOW_OFFLINE_DEDUP") == "1":
            return []
        return [f"无法读取历史交付完整性描述用于 G12 查重: {exc}"]
    errors = []
    for side, result in (review.get("sides") or {}).items():
        if result.get("decision") in {"EXACT", "SIMILAR"}:
            errors.append(f"G12 {side}-交付完整性描述与历史描述重复（{result.get('decision')}）：{result.get('rewriteInstruction') or ''}")
    return errors


def _validate_reason_quality(draft: dict[str, Any], evidence_doc: dict[str, Any]) -> list[str]:
    """提交前 preflight 的 reason-quality 门禁提前到草稿阶段（A/B 过程与产物、负面触发点和客观后果）。"""
    try:
        return list(_load_submission_preflight().assess_reason_quality(draft, evidence_doc).get("errors") or [])
    except Exception as exc:
        return [f"GSB 理由质量校验器不可用: {exc}"]


def _validate_claim_evidence(draft: dict[str, Any], evidence_doc: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    evidence_by_id = {
        str(item.get("id")): item
        for item in evidence_doc.get("evidence") or []
        if isinstance(item, dict) and item.get("id")
    }
    claims = draft.get("claims")
    if not isinstance(claims, list) or not claims:
        return ["draft.claims 必须是非空数组"]
    for side in ("A", "B"):
        side_claims = [item for item in claims if isinstance(item, dict) and item.get("side") == side]
        if not any(item.get("polarity") == "positive" for item in side_claims):
            errors.append(f"{side} 缺少正面证据声明")
        if not any(item.get("polarity") == "negative" for item in side_claims):
            errors.append(f"{side} 缺少负面证据声明")
        if not any(item.get("type") == "process" for item in side_claims):
            errors.append(f"{side} 缺少过程层面 claim，无法证明执行过程已覆盖")
        if not any(item.get("type") == "artifact" for item in side_claims):
            errors.append(f"{side} 缺少产物层面 claim，无法证明交付结果已覆盖")
    for index, claim in enumerate(claims, 1):
        if not isinstance(claim, dict):
            errors.append(f"claim {index} 不是对象")
            continue
        side = str(claim.get("side") or "")
        claim_type = str(claim.get("type") or "")
        polarity = str(claim.get("polarity") or "")
        if side not in {"A", "B"}:
            errors.append(f"claim {index} side 必须是 A/B")
        if claim_type not in {"process", "artifact"}:
            errors.append(f"claim {index} type 必须是 process/artifact")
        if polarity not in {"positive", "negative", "neutral"}:
            errors.append(f"claim {index} polarity 无效")
        evidence_ids = claim.get("evidenceIds") or []
        if not evidence_ids:
            errors.append(f"claim {index} 缺少 evidenceIds")
            continue
        for evidence_id in evidence_ids:
            evidence = evidence_by_id.get(str(evidence_id))
            if not evidence:
                errors.append(f"claim {index} 引用不存在的证据 {evidence_id}")
                continue
            if evidence.get("side") != side:
                errors.append(f"claim {index} 引用不同 side 的证据 {evidence_id}")
            if evidence.get("type") != claim_type:
                errors.append(f"claim {index} 引用类型不一致的证据 {evidence_id}")
            if _evaluation_excluded(evidence):
                errors.append(
                    f"claim {index} 不得引用 evaluationExcluded 的环境或工具噪声证据 {evidence_id}"
                )
    return errors


def _validate_negative_claim_triggers(draft: dict[str, Any], reason: str) -> list[str]:
    """Require every negative claim to locate the problem and state its objective consequence."""
    errors: list[str] = []
    claims = draft.get("claims")
    if not isinstance(claims, list):
        return errors
    clean_reason = _normalize(reason)
    if MACHINE_STEP_RE.search(reason):
        errors.append("GSB 理由出现三位以上的机器式步骤序号，请改写成用户可感知的触发节点")
    for index, claim in enumerate(claims, 1):
        if not isinstance(claim, dict) or claim.get("polarity") != "negative":
            continue
        kind = str(claim.get("triggerKind") or "").strip()
        trigger = str(claim.get("trigger") or "").strip()
        consequence = str(claim.get("objectiveConsequence") or "").strip()
        if kind not in TRIGGER_KINDS:
            errors.append(
                f"claim {index} 负面评价缺少有效 triggerKind，"
                "必须是 step/file/command/requirement"
            )
        if not trigger:
            errors.append(f"claim {index} 负面评价缺少 trigger 触发节点")
            continue
        clean_trigger = _normalize(trigger)
        if len(clean_trigger) < 4:
            errors.append(f"claim {index} trigger 过短，无法定位触发环节: {trigger}")
        if GENERIC_TRIGGER_RE.fullmatch(clean_trigger):
            errors.append(
                f"claim {index} trigger 只写了泛化过程“{trigger}”，"
                "必须落到具体步骤、文件、命令或需求"
            )
        if kind == "step" and not STEP_TRIGGER_RE.search(trigger):
            errors.append(f"claim {index} step trigger 缺少具体步骤或阶段: {trigger}")
        elif kind == "step" and NUMBERED_STEP_RE.search(trigger) and not STEP_ACTION_RE.search(trigger):
            errors.append(f"claim {index} step trigger 只有序号，没有说明执行了什么操作: {trigger}")
        elif kind == "file" and not FILE_TOKEN.search(trigger):
            errors.append(f"claim {index} file trigger 缺少文件名或路径: {trigger}")
        elif kind == "command" and not COMMAND_TRIGGER_RE.search(trigger):
            errors.append(f"claim {index} command trigger 缺少具体命令或工具: {trigger}")
        elif kind == "requirement" and not REQUIREMENT_TRIGGER_RE.search(trigger):
            errors.append(
                f"claim {index} requirement trigger 缺少触发需求或操作动作: {trigger}"
            )
        if clean_reason and clean_trigger not in clean_reason:
            errors.append(
                f"claim {index} 的 trigger 未原样写进 GSB 理由: {trigger}"
            )
        if not consequence:
            errors.append(f"claim {index} 负面评价缺少 objectiveConsequence 客观后果")
            continue
        clean_consequence = _normalize(consequence)
        if len(clean_consequence) < 4:
            errors.append(f"claim {index} objectiveConsequence 过短，无法说明实际影响: {consequence}")
        if not OBJECTIVE_CONSEQUENCE_RE.search(clean_consequence):
            errors.append(
                f"claim {index} objectiveConsequence 没有写出可见影响: {consequence}"
            )
        if clean_reason and clean_consequence not in clean_reason:
            errors.append(
                f"claim {index} 的 objectiveConsequence 未原样写进 GSB 理由: {consequence}"
            )
    return errors


def _validate_sentence_evidence(reason: str, draft: dict[str, Any], evidence_doc: dict[str, Any]) -> list[str]:
    """Require every reason sentence to map to existing evidence IDs."""
    errors: list[str] = []
    sentences = [part.strip() for part in re.split(r"(?<=[。；！？!?])", reason) if part.strip()]
    mappings = draft.get("sentenceEvidence")
    if not isinstance(mappings, list) or not mappings:
        return ["draft.sentenceEvidence 必须为每个理由句子提供证据映射"]
    evidence_by_id = {
        str(item.get("id")): item
        for item in evidence_doc.get("evidence") or []
        if isinstance(item, dict) and item.get("id")
    }
    evidence_ids = set(evidence_by_id)
    mapped_sentences: set[str] = set()
    for index, item in enumerate(mappings, 1):
        if not isinstance(item, dict):
            errors.append(f"sentenceEvidence {index} 不是对象")
            continue
        sentence = str(item.get("sentence") or "").strip()
        ids = item.get("evidenceIds") or []
        if not sentence:
            errors.append(f"sentenceEvidence {index} 缺少 sentence")
            continue
        if sentence not in sentences:
            errors.append(f"sentenceEvidence {index} 的句子不在 GSB 理由中: {sentence[:60]}")
        if not isinstance(ids, list) or not ids:
            errors.append(f"sentenceEvidence {index} 缺少 evidenceIds")
            continue
        for evidence_id in ids:
            if str(evidence_id) not in evidence_ids:
                errors.append(f"sentenceEvidence {index} 引用不存在的证据 {evidence_id}")
            elif _evaluation_excluded(evidence_by_id.get(str(evidence_id)) or {}):
                errors.append(
                    f"sentenceEvidence {index} 不得引用 evaluationExcluded 的环境或工具噪声证据 {evidence_id}"
                )
        mapped_sentences.add(sentence)
    for sentence in sentences:
        if sentence not in mapped_sentences:
            errors.append(f"GSB 理由句子缺少证据映射: {sentence[:60]}")
    return errors


def _validate_reason_layer_coverage(draft: dict[str, Any], reason: str) -> list[str]:
    """Require each side to put concrete process and artifact facts into the reason.

    claim.text is the reason anchor, not a private audit note.  A process fact must contain
    an actual action and a locatable file/command/step/requirement; an artifact fact must
    contain an observable product outcome.  This prevents a product-only paragraph from
    passing merely because it contains words such as "修改" or "测试".
    """
    errors: list[str] = []
    claims = draft.get("claims")
    if not isinstance(claims, list):
        return ["draft.claims 必须是非空数组"]
    clean_reason = _normalize(reason)
    labels = {"process": "过程", "artifact": "产物"}
    for side in ("A", "B"):
        side_claims = [item for item in claims if isinstance(item, dict) and item.get("side") == side]
        for claim_type in ("process", "artifact"):
            candidates = [item for item in side_claims if item.get("type") == claim_type]
            if not candidates:
                errors.append(f"{side} 侧缺少{labels[claim_type]} claim，GSB 理由必须分别覆盖 A/B 两层")
                continue
            selected = False
            for claim in candidates:
                claim_text = str(claim.get("text") or "").strip()
                clean_text = _normalize(claim_text)
                if len(clean_text) < 8 or clean_text not in clean_reason:
                    continue
                if claim_type == "process":
                    has_action = bool(PROCESS_ACTION_RE.search(claim_text))
                    has_locator = bool(PROCESS_LOCATOR_RE.search(claim_text) or FILE_TOKEN.search(claim_text))
                    if has_action and has_locator:
                        selected = True
                        break
                else:
                    if ARTIFACT_OUTCOME_RE.search(claim_text) or FILE_TOKEN.search(claim_text):
                        selected = True
                        break
            if not selected:
                if claim_type == "process":
                    errors.append(
                        f"{side} 侧过程 claim 未原样写入理由，或只有泛化过程；需同时写清读取/修改/执行等实际动作，"
                        "以及文件、命令、步骤、需求等定位点"
                    )
                else:
                    errors.append(
                        f"{side} 侧产物 claim 未原样写入理由，或没有可观察结果；需写清返回、缺少、未实现、失败、"
                        "写入、生成、通过等产物事实"
                    )
    return errors


def _validate_low_value_noise_claims(draft: dict[str, Any]) -> list[str]:
    """Allow noise only as reason context, never as its own evaluation claim."""
    errors: list[str] = []
    claims = draft.get("claims")
    if not isinstance(claims, list):
        return errors
    for index, claim in enumerate(claims, 1):
        if not isinstance(claim, dict):
            continue
        text = str(claim.get("text") or "")
        match = LOW_VALUE_PROCESS_NOISE_PATTERN.search(text)
        if match:
            errors.append(
                f"claim {index} 将低价值环境/工具噪声“{match.group(0)}”作为独立评价声明；"
                "这类信息只能作为真实业务或产物问题的原因描述，不能直接参与 A/B 评分或胜负判断"
            )
    return errors


def _validate_reason_markdown(reason: str) -> list[str]:
    """Reject every Markdown construct in the plain-text GSB reason."""
    errors: list[str] = []
    for label, pattern in MARKDOWN_REASON_PATTERNS:
        match = pattern.search(reason)
        if match:
            errors.append(
                f"GSB 理由不得包含 Markdown 语法（{label}）：{match.group(0)[:80]}"
            )
    return errors


def longest_common_fragment(left: str, right: str) -> int:
    """Length of the longest shared substring after whitespace normalization."""
    a, b = _normalize(left), _normalize(right)
    if not a or not b:
        return 0
    match = difflib.SequenceMatcher(None, a, b, autojunk=False).find_longest_match(0, len(a), 0, len(b))
    return match.size


def _delivery_label(side: str) -> str:
    return f"{side}-交付完整性描述"


def delivery_text_errors(side: str, description: str) -> list[str]:
    """Plain-text, wording and fluency rules shared with GSB 理由, relabelled per side."""
    label = _delivery_label(side)
    raw: list[str] = []
    raw.extend(_validate_reason_markdown(description))
    raw.extend(reason_style_errors(description))
    raw.extend(reason_language_errors(description))
    errors = [item.replace("GSB 理由", label) for item in raw]
    field = FIELD_FACTOR_RE.search(description)
    if field:
        errors.append(f"{label}不得引用录屏、截图、浏览器、运行环境等场外因素: {field.group(0)}")
    attribution = SOURCE_ATTRIBUTION_RE.search(description)
    if attribution:
        errors.append(
            f"{label}不要交代信息来源，直接写做了什么、看到什么，例如“请求了登录接口，返回404”: {attribution.group(0)}"
        )
    acceptance = NON_CONTAINER_TEST_ARTIFACT_RE.search(description) or DELIVERY_LOCAL_AUTOMATION_RE.search(description)
    if acceptance:
        errors.append(
            f"{label}不得引用本地（容器外）编写的测试或自动化脚本，这些不参与交付完整性描述: {acceptance.group(0)}"
        )
    return errors


def _side_trace_text(evidence_doc: dict[str, Any], side: str) -> str:
    """Raw JSONL of this side's trace; 平台拿本侧轨迹核验描述锚点。"""
    trace_path = str(((evidence_doc.get("process") or {}).get(side) or {}).get("tracePath") or "")
    if not trace_path:
        for item in evidence_doc.get("evidence") or []:
            if isinstance(item, dict) and item.get("side") == side and isinstance(item.get("trace"), dict):
                trace_path = str(item["trace"].get("tracePath") or "")
                if trace_path:
                    break
    path = Path(trace_path) if trace_path else None
    if path is None or not path.is_file():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def _is_recording(item: dict[str, Any]) -> bool:
    return str(item.get("id") or "").endswith("-recording")


def _recording_mode(item: dict[str, Any]) -> str:
    artifact = item.get("artifact") if isinstance(item.get("artifact"), dict) else {}
    return str(artifact.get("recordingMode") or "")


def delivery_excluded(item: dict[str, Any]) -> bool:
    """交付完整性描述不参考的证据：环境噪声、本地编写的测试/自动化脚本。

    有页面的项目（录屏 mode=web）和之前一样，录屏证据照常参与；
    纯后端 API 项目的录屏由本地 apiRequests 驱动，排除，改用验证计划里的 probe 接口探活。
    """
    artifact = item.get("artifact") if isinstance(item.get("artifact"), dict) else {}
    return (
        _evaluation_excluded(item)
        or bool(item.get("localScript") or artifact.get("localScript"))
        or (_is_recording(item) and _recording_mode(item) != "web")
    )


def _trace_file_activity(trace_text: str) -> tuple[set[str], str]:
    """Return (edited file paths via edit tools, all Bash command text) from raw Claude JSONL."""
    edited: set[str] = set()
    bash: list[str] = []
    for line in trace_text.splitlines():
        try:
            event = json.loads(line)
        except (json.JSONDecodeError, ValueError):
            continue
        content = ((event.get("message") or {}).get("content") if isinstance(event, dict) else None) or []
        if not isinstance(content, list):
            continue
        for block in content:
            if not isinstance(block, dict) or block.get("type") != "tool_use":
                continue
            tool_input = block.get("input") or {}
            if block.get("name") in EDIT_TOOL_NAMES:
                path = str(tool_input.get("file_path") or tool_input.get("notebook_path") or "")
                if path:
                    edited.add(path)
            elif block.get("name") == "Bash":
                bash.append(str(tool_input.get("command") or ""))
    return edited, "\n".join(bash)


def _edit_contradiction_errors(label: str, description: str, trace_text: str) -> list[str]:
    edited, bash = _trace_file_activity(trace_text)
    errors: list[str] = []
    for sentence in _reason_sentences(description):
        for clause in re.split(r"[，,；;]", sentence):
            for token in set(FILE_TOKEN.findall(clause)):
                name = Path(token).name
                touched = any(path.endswith(token) or Path(path).name == name for path in edited)
                if DELIVERY_NOT_CHANGED_RE.search(clause) and touched:
                    errors.append(f"{label}说 {token} 没有改，但本侧轨迹里有对它的编辑，与轨迹完全对立")
                elif DELIVERY_CHANGED_RE.search(clause) and not touched and name not in bash:
                    errors.append(f"{label}说改了 {token}，但本侧轨迹里没有任何对它的编辑，与轨迹完全对立")
    return errors


def _is_runnability_check(item: dict[str, Any]) -> bool:
    artifact = item.get("artifact") or {}
    text = " ".join(str(value or "") for value in (item.get("text"), artifact.get("command")))
    return bool(DELIVERY_RUNNABILITY_RE.search(text))


def _evidence_failed(item: dict[str, Any]) -> bool:
    artifact = item.get("artifact") or {}
    if isinstance(artifact, dict) and ("ok" in artifact or "observedFailure" in artifact):
        return bool(artifact.get("observedFailure")) or not bool(artifact.get("ok", True))
    return item.get("polarity") == "negative"


def delivery_trace_consistency_errors(
    side: str,
    score: int,
    description: str,
    cited_ids: list[Any],
    draft: dict[str, Any],
    evidence_doc: dict[str, Any],
) -> list[str]:
    """红线：交付完整性描述必须与本侧轨迹、真实复核和 GSB 理由一致，不得出现对立意见。"""
    label = _delivery_label(side)
    errors: list[str] = []
    evidence = [
        item for item in evidence_doc.get("evidence") or []
        if isinstance(item, dict) and item.get("side") == side and not delivery_excluded(item)
    ]
    by_id = {str(item.get("id")): item for item in evidence}

    # 锚点：文件、命令、报错必须出现在本侧原始轨迹；状态码与接口路径可来自本侧真实复核输出。
    trace_text = _side_trace_text(evidence_doc, side)
    if not trace_text:
        errors.append(f"{label}无法读取本侧轨迹文件，不能核对描述与轨迹是否一致")
    else:
        artifact_text = json.dumps(
            [item for item in evidence if item.get("type") == "artifact"], ensure_ascii=False
        )
        for token in sorted(set(FILE_TOKEN.findall(description))):
            if token not in trace_text and Path(token).name not in trace_text:
                errors.append(f"{label}提到的文件 {token} 在本侧轨迹中不存在")
        for token in sorted(set(DELIVERY_LITERAL_ANCHOR_RE.findall(description))):
            if token not in trace_text:
                errors.append(f"{label}提到的命令或报错 {token} 在本侧轨迹中不存在")
        for token in sorted(set(DELIVERY_OBSERVED_ANCHOR_RE.findall(description))):
            if token not in trace_text and token not in artifact_text:
                errors.append(f"{label}提到的 {token} 在本侧轨迹和真实复核输出中都找不到")
        errors.extend(_edit_contradiction_errors(label, description, trace_text))

    # 与真实复核结果对立。
    checks = [item for item in evidence if item.get("type") == "artifact" and item.get("artifact")]
    failed = [item for item in checks if _evidence_failed(item)]
    run_failed = [item for item in failed if _is_runnability_check(item)]
    if run_failed and score >= 3:
        errors.append(
            f"{label}给 {score} 分，但本侧构建或启动复核失败（{run_failed[0].get('id')}）；"
            "无法运行最高 2 分"
        )
    if failed and score == 5:
        errors.append(f"{label}给 5 分，但本侧复核存在失败项（{failed[0].get('id')}），与“一次性完整跑通”对立")
    if run_failed:
        claim = DELIVERY_RUN_OK_RE.search(description)
        if claim:
            errors.append(f"{label}写了“{claim.group(0)}”，但本侧构建或启动复核失败，描述与复核结果对立")
    if checks and not failed:
        if score == 1:
            errors.append(f"{label}给 1 分，但本侧复核全部通过，与“完全失败”对立")
        claim = DELIVERY_RUN_FAIL_RE.search(description)
        if claim:
            errors.append(f"{label}写了“{claim.group(0)}”，但本侧构建、启动与复核全部通过，描述与复核结果对立")

    # 引用证据的正负方向必须与分数一致。
    cited = [by_id[str(item)] for item in cited_ids if str(item) in by_id]
    has_pages = any(_is_recording(item) and _recording_mode(item) == "web" for item in evidence)
    if has_pages and any(((item.get("artifact") or {}).get("probe")) for item in cited):
        errors.append(f"{label}所在项目有页面，不做接口探活；按页面操作结果引用本侧录屏证据")
    if score and score < 5 and not any(_evidence_failed(item) for item in cited):
        if failed:
            errors.append(f"{label}不给 5 分，但引用的本侧证据全是正面结果；至少引用一条能证明问题的失败证据")
        else:
            errors.append(
                f"{label}不给 5 分，但本侧没有任何可引用的失败证据；纯后端 API 项目在验证计划里"
                "给这一侧补一条启动加 probe 接口探活后重新 verify，有页面的项目引用本侧录屏证据"
            )
    if score == 5 and any(_evidence_failed(item) for item in cited):
        errors.append(f"{label}给 5 分却引用了本侧失败证据，分数与证据对立")

    # 与 GSB 理由里本侧的产物结论对立。
    negative_artifact = [
        claim for claim in draft.get("claims") or []
        if isinstance(claim, dict) and claim.get("side") == side
        and claim.get("type") == "artifact" and claim.get("polarity") == "negative"
    ]
    if score == 5 and negative_artifact:
        errors.append(
            f"{label}给 5 分，但 GSB 理由里本侧有产物负面结论“{str(negative_artifact[0].get('text') or '')[:40]}”，两处意见对立"
        )

    # 虚假成功必须以轨迹里模型的最终宣称为依据。
    if "虚假成功" in description:
        final = next((item for item in evidence if item.get("id") == f"{side}-process-final"), None)
        quote = str(((final or {}).get("trace") or {}).get("quote") or "")
        if not DELIVERY_COMPLETION_CLAIM_RE.search(quote):
            errors.append(f"{label}判定虚假成功，但本侧轨迹的最终回复里没有宣称完成，缺少依据")
    return errors


def validate_delivery(
    draft: dict[str, Any],
    evidence_doc: dict[str, Any],
    reason: str,
) -> dict[str, Any]:
    """Validate A/B 交付完整性打分与描述（官方字段 a/b_score_delivery、a/b_desc_delivery）。

    评分锚点见 references/delivery-scoring.md；描述只谈交付完整性，
    允许与 GSB 理由有少量重合，但不得照抄理由，A/B 两段之间也不得雷同（平台规则 G12）。
    """
    errors: list[str] = []
    warnings: list[str] = []
    delivery = draft.get("delivery")
    if not isinstance(delivery, dict):
        return {"ok": False, "errors": ["draft.delivery 必填：需要 A、B 两侧的 score 与 description"], "warnings": []}
    evidence_by_id = {
        str(item.get("id")): item
        for item in evidence_doc.get("evidence") or []
        if isinstance(item, dict) and item.get("id")
    }
    scores: dict[str, int] = {}
    descriptions: dict[str, str] = {}
    for side in ("A", "B"):
        label = _delivery_label(side)
        entry = delivery.get(side)
        if not isinstance(entry, dict):
            errors.append(f"draft.delivery.{side} 缺失")
            continue
        score = entry.get("score")
        if isinstance(score, bool) or not isinstance(score, int) or not 1 <= score <= 5:
            errors.append(f"{side}-交付完整性必须是 1~5 的整数，当前 {score!r}")
            score = 0
        description = str(entry.get("description") or "").strip()
        length = text_non_whitespace_len(description)
        if not DELIVERY_DESC_MIN <= length <= DELIVERY_DESC_MAX:
            errors.append(
                f"{label}必须为 {DELIVERY_DESC_MIN}–{DELIVERY_DESC_MAX} 个非空白字符，当前 {length}"
            )
        if not description:
            continue
        scores[side] = score
        descriptions[side] = description
        errors.extend(delivery_text_errors(side, description))
        dimension = DELIVERY_PROCESS_DIMENSION_RE.search(description)
        if dimension:
            errors.append(
                f"{label}只评价交付完整性，规划、推理、工具调用等过程维度写进 GSB 理由: {dimension.group(0)}"
            )
        if score == 5:
            if not DELIVERY_VERIFY_RE.search(description):
                errors.append(f"{label}给 5 分必须写出核对依据：逐条核对了哪些需求、跑过什么验证及结论")
            unresolved = DELIVERY_UNRESOLVED_RE.search(description)
            if unresolved:
                errors.append(f"{label}给 5 分却写了未解决问题“{unresolved.group(0)}”，分数与描述矛盾")
        elif score:
            if not (OBJECTIVE_CONSEQUENCE_RE.search(description) or DELIVERY_CONSEQUENCE_RE.search(description)):
                errors.append(f"{label}不给 5 分必须写出造成的客观后果，例如没有写入、返回404、无法启动")
            if not (FILE_TOKEN.search(description) or PROCESS_LOCATOR_RE.search(description) or ERROR_TOKEN.search(description)):
                errors.append(f"{label}不给 5 分必须写清问题出在哪：文件名、报错原文或未实现的需求点")
        if "虚假成功" in description and not (
            DELIVERY_CLAIMED_RE.search(description) and DELIVERY_ACTUAL_RE.search(description)
        ):
            errors.append(f"{label}判定虚假成功时，要写清模型宣称改了什么、实际改了什么")
        ids = entry.get("evidenceIds") or []
        if not isinstance(ids, list) or not ids:
            errors.append(f"draft.delivery.{side}.evidenceIds 必填，描述要能对回本侧证据")
            ids = []
        for evidence_id in ids:
            evidence = evidence_by_id.get(str(evidence_id))
            if not evidence:
                errors.append(f"draft.delivery.{side} 引用不存在的证据 {evidence_id}")
            elif evidence.get("side") != side:
                errors.append(f"draft.delivery.{side} 引用了另一侧的证据 {evidence_id}")
            elif _evaluation_excluded(evidence):
                errors.append(f"draft.delivery.{side} 不得引用 evaluationExcluded 证据 {evidence_id}")
            elif delivery_excluded(evidence):
                errors.append(
                    f"draft.delivery.{side} 不得引用本地编写的测试、自动化脚本或录制脚本产生的证据 {evidence_id}"
                )
        errors.extend(
            delivery_trace_consistency_errors(side, score, description, ids, draft, evidence_doc)
        )
        copied = longest_common_fragment(description, reason)
        if copied >= DELIVERY_REASON_COPY_FRAGMENT:
            errors.append(
                f"{label}与 GSB 理由有 {copied} 字连续相同，属于照抄；只从完整性角度重新组织"
            )
        elif difflib.SequenceMatcher(None, _normalize(description), _normalize(reason)).ratio() >= DELIVERY_REASON_COPY_RATIO:
            errors.append(f"{label}与 GSB 理由整体过于相似，只从完整性角度重新组织")
    if len(descriptions) == 2:
        shared = longest_common_fragment(descriptions["A"], descriptions["B"])
        ratio = difflib.SequenceMatcher(
            None, _normalize(descriptions["A"]), _normalize(descriptions["B"])
        ).ratio()
        if shared >= DELIVERY_SIDE_FRAGMENT or ratio >= DELIVERY_SIDE_RATIO:
            errors.append(
                f"A、B 两段交付完整性描述雷同（连续相同 {shared} 字，相似度 {ratio:.0%}）；"
                "平台 G12 会拿两段互相比对，要按各自实际情况独立写"
            )
    verdict = str(draft.get("verdict") or "")
    if len(scores) == 2 and all(scores.values()):
        # 红线：结论与分数方向不得对立。
        if verdict == "A 更好" and scores["A"] < scores["B"]:
            errors.append("GSB 结论为 A 更好，但 A 的交付完整性低于 B，结论与打分对立")
        if verdict == "B 更好" and scores["B"] < scores["A"]:
            errors.append("GSB 结论为 B 更好，但 B 的交付完整性低于 A，结论与打分对立")
        if verdict == "Same" and abs(scores["A"] - scores["B"]) >= 2:
            errors.append("GSB 结论为 Same，但两侧交付完整性相差 2 分以上，结论与打分对立")
    errors = list(dict.fromkeys(errors))
    return {"ok": not errors, "errors": errors, "warnings": warnings, "scores": scores}


def validate_draft(draft: dict[str, Any], task_root: Path, *, review_path: Path) -> dict[str, Any]:
    errors: list[str] = []
    evidence_doc = read_json(task_root / "monitor" / "evidence.json", {})
    if not evidence_doc.get("evidence"):
        errors.append("缺少证据索引，请先 audit")
    verdict = str(draft.get("verdict") or "")
    if verdict not in {"A 更好", "Same", "B 更好"}:
        errors.append("verdict 必须是 A 更好、Same、B 更好")
    reason = str(draft.get("reason") or "").strip()
    # 2026-09-23 官方表单已删除“备注”字段；旧草稿残留的 remark 直接丢弃。
    draft.pop("remark", None)
    length = text_non_whitespace_len(reason)
    if not 150 <= length <= 240:
        errors.append(f"GSB 理由必须为 150–240 个非空白字符，当前 {length}")
    markdown_errors = _validate_reason_markdown(reason)
    language_errors = reason_language_errors(reason)
    excluded_errors = evaluation_excluded_reason_errors(reason, draft, evidence_doc)
    errors.extend(markdown_errors)
    errors.extend(language_errors)
    flow_errors = reason_flow_errors(reason)
    errors.extend(flow_errors)
    errors.extend(_validate_artifact_description(reason))
    for side, full_label in (("A", "A 侧方案"), ("B", "B 侧方案")):
        if full_label not in reason:
            errors.append(f"GSB 理由必须使用完整表述“{full_label}”，禁止省略式单字")
        if not re.search(rf"\b{side}\b|[（(]{side}[）)]|{side}侧|{side}的表现|{side}跑", reason):
            errors.append(f"GSB 理由没有明确覆盖 {side}")
    if verdict == "Same" and not re.search(r"Same|等价|抵消|持平|无明显差异|难分高下", reason):
        errors.append("Same 必须写明等价点和相互抵消项")
    errors.extend(_validate_reason_layer_coverage(draft, reason))
    errors.extend(_validate_sentence_evidence(reason, draft, evidence_doc))
    errors.extend(_validate_claim_evidence(draft, evidence_doc))
    errors.extend(_validate_low_value_noise_claims(draft))
    errors.extend(excluded_errors)
    errors.extend(_validate_negative_claim_triggers(draft, reason))
    evidence_text = json.dumps(evidence_doc, ensure_ascii=False)
    for token in sorted(set(FILE_TOKEN.findall(reason))):
        if token not in evidence_text:
            errors.append(f"理由中的文件 {token} 没有证据")
    for token in sorted(set(ERROR_TOKEN.findall(reason))):
        token = token.strip()
        if len(token) > 4 and token not in evidence_text:
            # Error wording may be paraphrased; require the first error keyword plus a literal excerpt.
            literal_parts = re.findall(r"[A-Za-z][A-Za-z0-9_.:/!-]{4,}|[\u4e00-\u9fff]{4,}", token)
            if literal_parts and not any(part in evidence_text for part in literal_parts):
                errors.append(f"理由中的报错描述没有证据: {token[:80]}")
    numeric_tokens = re.findall(r"(?<![A-Za-z0-9])\d+(?:\.\d+)?(?![A-Za-z0-9])", reason)
    if len(numeric_tokens) > 2:
        errors.append(f"GSB 理由数字对比过多，当前 {len(numeric_tokens)} 处；只保留必要的少量事实")
    test_count_match = TEST_COUNT_PATTERN.search(reason)
    if test_count_match:
        errors.append(
            "GSB 理由不要用测试或断言数量堆砌，改用业务覆盖描述："
            f"{test_count_match.group(0)}"
        )
    file_tokens = sorted(set(FILE_TOKEN.findall(reason)))
    if len(file_tokens) > 2:
        errors.append(f"GSB 理由代码细节过多，当前 {len(file_tokens)} 个文件/代码标记")
    errors.extend(_validate_reason_dedup(reason, task_root))
    errors.extend(_validate_delivery_dedup(draft, task_root))
    errors.extend(_validate_reason_quality(draft, evidence_doc))
    delivery = validate_delivery(draft, evidence_doc, reason)
    errors.extend(delivery["errors"])

    review_errors: list[str] = []
    if not AUDIT_HUMAN.is_file():
        review_errors.append(f"缺少 ra-人话审核器: {AUDIT_HUMAN}")
    else:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", delete=False) as temp:
            temp.write(reason)
            temp_path = Path(temp.name)
        try:
            proc = subprocess.run(
                [
                    sys.executable,
                    str(AUDIT_HUMAN),
                    "--text-file",
                    str(temp_path),
                    "--target-type",
                    "comments",
                    "--review",
                    str(review_path),
                    "--json",
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            try:
                result = json.loads(proc.stdout)
                review_errors.extend(result.get("errors") or [])
            except json.JSONDecodeError:
                review_errors.append(proc.stderr.strip() or proc.stdout.strip() or "ra-人话审核器无 JSON 输出")
        finally:
            temp_path.unlink(missing_ok=True)
    errors.extend(review_errors)
    errors = list(dict.fromkeys(error for error in errors if error))
    return {
        "ok": not errors,
        "verdict": verdict,
        "reason": reason,
        "reasonLength": length,
        "reasonSha256": sha256_text(reason),
        "markdownErrors": markdown_errors,
        "languageErrors": language_errors,
        "evaluationExcludedErrors": excluded_errors,
        "flowErrors": flow_errors,
        "styleWarnings": reason_style_warnings(reason),
        "delivery": delivery,
        "errors": errors,
    }


def _video_expected(task_root: Path, side: str) -> Path:
    return task_root / "workspace" / "视频信息" / side.lower() / "视频" / "demo.mp4"


def canonical_draft_path(task_root: Path) -> Path:
    """提交预检固定读取的草稿位置（monitor/gsb-draft.json）。"""
    return task_root / "monitor" / "gsb-draft.json"


def sync_canonical_draft(task_root: Path, values: dict[str, Any]) -> Path:
    """把 GSB 草稿的关键字段同步到提交预检读取的固定路径。

    预检同时兼容任务根目录的 gsb-draft.json；这里只维护 monitor/gsb-draft.json，
    保留既有草稿里的 claims / sentenceEvidence 等审计字段。
    """
    path = canonical_draft_path(task_root)
    draft = read_json(path, {}) if path.is_file() else {}
    if not isinstance(draft, dict):
        draft = {}
    draft.setdefault("schemaVersion", 1)
    draft.update(
        {
            "verdict": str(values.get("gsb_verdict") or ""),
            "reason": str(values.get("gsb_reason") or "").strip(),
            "languages": str(values.get("languages") or ""),
            "repro_level": str(values.get("repro_level") or ""),
        }
    )
    draft.pop("remark", None)
    delivery = draft.get("delivery") if isinstance(draft.get("delivery"), dict) else {}
    for side in ("A", "B"):
        prefix = side.lower()
        entry = dict(delivery.get(side) or {})
        if str(values.get(f"{prefix}_score_delivery") or "").strip():
            entry["score"] = int(values[f"{prefix}_score_delivery"])
        if f"{prefix}_desc_delivery" in values:
            entry["description"] = str(values[f"{prefix}_desc_delivery"] or "").strip()
        if entry:
            delivery[side] = entry
    if delivery:
        draft["delivery"] = delivery
    write_json(path, draft)
    return path


def build_values(task_root: Path, draft: dict[str, Any], schema: dict[str, Any]) -> dict[str, Any]:
    state = read_json(task_root / "monitor" / "state.json", {})
    side_a = (state.get("sides") or {}).get("A") or {}
    side_b = (state.get("sides") or {}).get("B") or {}
    raw_harness_version = str(side_a.get("harnessVersion") or side_b.get("harnessVersion") or "").strip()
    harness_match = re.search(r"(?<!\d)(\d+(?:\.\d+)+)", raw_harness_version)
    harness_version = harness_match.group(1) if harness_match else raw_harness_version
    trace_a = ensure_single_side_trace(task_root, "A", side_a.get("tracePath"))
    trace_b = ensure_single_side_trace(task_root, "B", side_b.get("tracePath"))
    recording_a = ((state.get("recordings") or {}).get("A") or {})
    recording_b = ((state.get("recordings") or {}).get("B") or {})
    video_a = Path(str(recording_a.get("videoPath") or _video_expected(task_root, "A")))
    video_b = Path(str(recording_b.get("videoPath") or _video_expected(task_root, "B")))
    if trace_a.stat().st_size > 27 * 1024 * 1024 or trace_b.stat().st_size > 27 * 1024 * 1024:
        raise SologsbError("轨迹文件超过 27MB 上限")
    if "languages" not in draft or not str(draft.get("languages") or "").strip():
        raise SologsbError("draft.languages 必填")
    if "repro_level" not in draft or not str(draft.get("repro_level") or "").strip():
        raise SologsbError("draft.repro_level 必填")
    delivery = draft.get("delivery") if isinstance(draft.get("delivery"), dict) else {}
    delivery_a = delivery.get("A") or {}
    delivery_b = delivery.get("B") or {}
    values = {
        "user_prompt": Path(str(state["promptPath"])).read_text(encoding="utf-8"),
        "question_type": state.get("taskType", ""),
        "difficulty": state.get("difficulty", ""),
        "languages": draft.get("languages", ""),
        "harness": "Claude Code",
        "harness_version": harness_version,
        "os_platform": "MacOS/Linux",
        "repro_level": draft.get("repro_level", ""),
        "env_snapshot": commit_url(str(state["repoUrl"]), str(state["initialSnapshot"])),
        "a_session_id": side_a.get("sessionId", ""),
        "a_trace_file": str(trace_a.resolve()),
        "a_artifact_snapshot": side_a.get("artifactSnapshotUrl", ""),
        "a_screencast": str(video_a.resolve()),
        "a_score_delivery": delivery_a.get("score", ""),
        "a_desc_delivery": str(delivery_a.get("description") or "").strip(),
        "b_session_id": side_b.get("sessionId", ""),
        "b_trace_file": str(trace_b.resolve()),
        "b_artifact_snapshot": side_b.get("artifactSnapshotUrl", ""),
        "b_screencast": str(video_b.resolve()),
        "b_score_delivery": delivery_b.get("score", ""),
        "b_desc_delivery": str(delivery_b.get("description") or "").strip(),
        "gsb_verdict": draft.get("verdict", ""),
        "gsb_reason": draft.get("reason", ""),
    }
    allowed = {str(field.get("field_key")): field for field in schema.get("fields") or []}
    unknown = set(values) - set(allowed)
    if unknown:
        raise SologsbError(f"内部值映射包含未知字段: {sorted(unknown)}")
    for key, field in allowed.items():
        if field.get("is_required") and not str(values.get(key) or "").strip():
            raise SologsbError(f"必填字段为空: {key}")
        rule = field.get("validation") or {}
        if field.get("field_type") == "number" and key in values:
            value = values[key]
            if isinstance(value, bool) or not isinstance(value, int):
                raise SologsbError(f"{key} 必须是整数，当前 {value!r}")
            if not int(rule.get("min", value)) <= value <= int(rule.get("max", value)):
                raise SologsbError(f"{key} 超出 {rule.get('min')}~{rule.get('max')}: {value}")
    return values


def write_excel(task_root: Path, schema: dict[str, Any], values: dict[str, Any]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "GSB提交"
    fields = sorted(schema.get("fields") or [], key=lambda item: int(item.get("sort_order") or 0))
    headers = [str(field.get("label") or field.get("field_key")) for field in fields]
    sheet.append(headers)
    sheet.append([values.get(str(field.get("field_key")), "") for field in fields])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(vertical="center")
    widths = [28, 28, 24, 24, 18, 24, 20, 30, 52, 42, 58, 52, 58, 16, 60, 42, 58, 52, 58, 16, 60, 16, 90]
    for index, width in enumerate(widths[: len(headers)], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    meta = workbook.create_sheet("schema")
    meta.append(["fingerprint", schema.get("fingerprint", "")])
    meta.append(["attachment_max_mb", schema.get("attachment_max_mb", 27)])
    meta.append(["video_max_mb", schema.get("video_max_mb", 500)])
    meta.append(["exported_at", utc_now()])
    output = task_root / "workspace" / "评审文件" / "交付表.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def write_field_guide(task_root: Path, schema: dict[str, Any], values: dict[str, Any]) -> Path:
    fields = sorted(schema.get("fields") or [], key=lambda item: int(item.get("sort_order") or 0))
    lines = [
        "# GSB 提交字段填写说明",
        "",
        f"- 官方 schema fingerprint：`{schema.get('fingerprint', '')}`",
        f"- 附件上限：{schema.get('attachment_max_mb', 27)} MB",
        f"- 视频上限：{schema.get('video_max_mb', 500)} MB",
        "- 本地阶段轨迹和录屏填写绝对路径；提交脚本会读取该路径并通过 API 上传。",
        "",
        "| 顺序 | 分组 | 字段 | 类型 | 必填 | 当前值/填写规则 |",
        "|---:|---|---|---|---|---|",
    ]
    for index, field in enumerate(fields, 1):
        key = str(field.get("field_key") or "")
        value = str(values.get(key) or "")
        if len(value) > 700:
            value = value[:700] + "…"
        value = value.replace("|", "\\|").replace("\n", "<br>")
        required = "是" if field.get("is_required") else "否"
        options = " / ".join(str(item) for item in field.get("options") or [])
        help_text = str(field.get("help_text") or "").replace("|", "\\|").replace("\n", "<br>")
        if options:
            value = f"{value}<br>枚举：{options}"
        if help_text:
            value = f"{value}<br>官方说明：{help_text}"
        if field.get("field_type") in {"attachment", "video"}:
            value = f"{value}<br>提交：`submit_api.py` 会上传文件，再把平台返回的对象地址写入 API payload；Excel 仍保留本地绝对路径。"
        if key in {"a_screencast", "b_screencast"}:
            value = f"{value}<br>视频规格：1280x720（720p），单段不超过 90 秒；Web 仅 Terminal.app+Chrome，终端/失败仅 Terminal.app。"
        lines.append(
            f"| {index} | {field.get('group', '')} | `{key}` / {field.get('label', '')} | "
            f"{field.get('field_type', '')} | {required} | {value} |"
        )
    lines.extend(
        [
            "",
            "## 规则提醒",
            "",
            "- `User Prompt` 与两份轨迹中的真人输入必须逐字一致。",
            "- A/B SessionID 必须不同，产物快照父提交必须同为初始环境快照。",
            "- Harness 版本字段只填数字版本，例如 `2.1.197`，不要带客户端名称后缀。",
            "- GSB 理由总字段必须为 150–240 个非空白字符，A/B 两侧都要分别覆盖过程与产物。",
            "- GSB 理由必须严格使用纯文本，不允许任何 Markdown 语法；标题、列表、代码块、行内代码、链接、图片、强调标记、表格、引用和 HTML 标签全部阻断。",
            "- 过程层写轨迹中的实际动作和定位节点，例如读取、检查、修改、执行、排查或返工了哪一步、文件、命令或需求；不能只写“进行了测试”或“做了迁移”。",
            "- 产物层写最终可观察结果，例如接口返回、缺少字段、未实现需求、构建启动结果或真实失败；不得只写交付物毛病。",
            "- 录屏、视频、截图、浏览器、测试设备、运行环境、验收宿主、Otty、Terminal.app、鼠标、分辨率等场外因素不得写进 GSB 理由，没有例外。",
            "- evaluationExcluded 的环境或工具噪声证据不得被 claim、sentenceEvidence 或理由正文引用；其他噪声也不能成为独立评分项或 A/B 胜负依据。",
            "- 用于支撑每侧过程/产物的 claim.text 必须原样出现在 GSB 理由中。",
            "- 每条负面 claim 必须提供 triggerKind/trigger；触发节点只允许步骤、文件、命令或需求，并原样写入 GSB 理由。",
            "- 轨迹超过 27MB 或视频超过 500MB 时，本地交付不得标记完成。",
            "- G10：理由与当前 GSB 数据精确重复或存在长连续片段复用时必须重写。",
            "- A/B-交付完整性：1~5 的整数，按 references/delivery-scoring.md 的五档锚点打分，两侧各自独立评分。",
            "- A/B-交付完整性描述：40–200 个非空白字符，只写需求是否做完、代码能否运行、有没有虚假成功；不写规划、推理、工具调用等过程维度。",
            "- 给 5 分写核对依据（核对了哪些需求、跑过什么验证、结论如何）；不给 5 分写清问题出在哪、模型做了什么、造成了什么客观后果。",
            "- 交付完整性描述允许与 GSB 理由有少量重合，但不得照抄（连续 20 字相同即阻断）；A、B 两段之间也不得雷同（G12）。",
            "- 红线：描述里提到的文件名、命令或报错必须在本侧原始轨迹中存在；分数与描述不得与本侧真实复核、引用证据、GSB 理由或 GSB 结论对立。",
            "- 本地（容器外）编写的任何测试和自动化脚本（验收、冒烟、Playwright、录制场景）不参与交付完整性描述，不写入正文、不作为引用证据。",
            "- 有页面的项目和之前一样引用录屏证据；纯后端 API 项目录屏不引用，用验证计划的 probe 接口探活证明问题。",
            "- 描述与理由都直接写“请求了登录接口，返回404”，不写“从录屏来看”“根据编写的测试”“复核结果显示”。",
            "- GSB 理由使用完整、质朴的中文描述，统一写“A 侧方案”“B 侧方案”，不使用省略式单字；每侧称谓最多出现 3 次，相邻两句不要用同一称谓起头，不写 8 字以下的碎句。",
            "- 不写电报体：全文最多 6 句，同一侧的动作和结果用逗号连成一句；结论句之前至少 2 处“结果”“导致”“但”“却”这类因果或转折衔接，让读者一眼看出谁因谁果、哪条是扣分点。",
            "- 不把两件事压成“先提交标题未更新”这类短语，补上主语和衔接，写成“先提交的一方写进去了，标题却没有更新”。",
            "- 结尾前用一句话交代这个任务最重要的是哪一条，再给结论；不要只罗列事实后直接宣布胜负。",
            "- 禁用“闭环”“根因”“落库”，也禁用旧句式“这题”“最要紧”，统一写“这个任务最重要的是”；数据写入统一写“入库”；常用命令“npm run build”统一写“build”，避免历史长片段去重；“真实”“真正”“其实”等空泛表达会给出警告。",
            "- 句子达到高中语文阅读水平，表达通顺；单句非空白字符不得超过 56 字，分句和标点异常会阻断。",
            "- 禁止使用“落在……”式收束句式；结论直接写“因此选择 B 侧方案”或“B 侧方案更好”。",
        ]
    )
    output = task_root / "workspace" / "评审文件" / "GSB提交字段说明.md"
    atomic_write_text(output, "\n".join(lines) + "\n")
    return output


def export_gsb(
    task_root: Path,
    *,
    draft_path: Path,
    review_path: Path,
    accept_schema_change: bool = False,
) -> dict[str, Any]:
    state = read_json(task_root / "monitor" / "state.json", {})
    if state.get("status") not in {"verified", "gsb_ready", "recorded", "complete"}:
        raise SologsbError("必须先进入 verified")
    run_audit(task_root)
    draft = read_json(draft_path)
    if not isinstance(draft, dict):
        raise SologsbError("GSB draft 必须是 JSON 对象")
    validation = validate_draft(draft, task_root, review_path=review_path)
    write_json(draft_path, draft)
    write_json(canonical_draft_path(task_root), draft)
    monitor_dir = task_root / "monitor" / "gsb"
    monitor_dir.mkdir(parents=True, exist_ok=True)
    write_json(monitor_dir / "validation.json", validation)
    if not validation["ok"]:
        raise SologsbError("GSB 门禁未通过:\n- " + "\n- ".join(validation["errors"]))
    schema = load_official_schema(accept_change=accept_schema_change)
    values = build_values(task_root, draft, schema)
    excel = write_excel(task_root, schema, values)
    guide = write_field_guide(task_root, schema, values)
    write_json(monitor_dir / "values.json", values)
    write_json(
        monitor_dir / "gsb.json",
        {
            "schemaVersion": 1,
            "verdict": draft.get("verdict"),
            "reason": draft.get("reason"),
            "reasonLength": validation["reasonLength"],
            "reasonSha256": validation["reasonSha256"],
            "reviewPath": str(review_path.resolve()),
            "excelPath": str(excel.resolve()),
            "fieldGuidePath": str(guide.resolve()),
        },
    )
    state["status"] = "gsb_ready"
    state["excelPath"] = str(excel.resolve())
    state["fieldGuidePath"] = str(guide.resolve())
    save_state(task_root, state)
    return {
        "status": "gsb_ready",
        "verdict": draft.get("verdict"),
        "reasonLength": validation["reasonLength"],
        "excelPath": str(excel.resolve()),
        "fieldGuidePath": str(guide.resolve()),
    }


def refresh_excel(task_root: Path, *, accept_schema_change: bool = False) -> Path:
    draft = read_json(task_root / "monitor" / "gsb" / "values.json", {})
    if not draft:
        raise SologsbError("缺少 GSB values.json")
    schema = load_official_schema(accept_change=accept_schema_change)
    # refresh video paths and harness version from current state
    state = read_json(task_root / "monitor" / "state.json", {})
    side_a = (state.get("sides") or {}).get("A") or {}
    side_b = (state.get("sides") or {}).get("B") or {}
    raw_harness = str(side_a.get("harnessVersion") or side_b.get("harnessVersion") or draft.get("harness_version", ""))
    harness_match = re.search(r"(?<!\d)(\d+(?:\.\d+)+)", raw_harness)
    draft["harness_version"] = harness_match.group(1) if harness_match else raw_harness
    for side, side_state in (("a", side_a), ("b", side_b)):
        recording = ((state.get("recordings") or {}).get(side.upper()) or {})
        video = Path(str(recording.get("videoPath") or _video_expected(task_root, side.upper())))
        if video.is_file():
            draft[f"{side}_screencast"] = str(video.resolve())
    excel = write_excel(task_root, schema, draft)
    write_field_guide(task_root, schema, draft)
    sync_canonical_draft(task_root, draft)
    return excel
