#!/usr/bin/env python3
"""CLI for the sologsb-0917 Pair-wise GSB skill."""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import re
from pathlib import Path
from typing import Any

# --- 设备配置注入：必须在下列模块读取环境变量之前执行 ---
sys.path.insert(0, str(Path(__file__).resolve().parent))
import device_config as _device_config  # noqa: E402

_device_config.load_and_apply()

from artifact_verifier import run_verification
from common import (
    DEFAULT_RECORDING_LOCK_TIMEOUT,
    SCHEMA_FALLBACK,
    SologsbError,
    ensure_task_dirs,
    github_env,
    load_state,
    read_json,
    run,
    safe_local_name,
    safe_slug,
    save_state,
    sha256_file,
    task_root_from_arg,
    utc_now,
    write_json,
)
from evidence import run_audit
from github_repo import init_github_repo
from gsb_tools import export_gsb
from prompt_tools import install_prompt
from project_claims import release_project_claim
from recorder import prepare_recording, record_side, recording_isolation_ok, video_dimensions
from semantic_review import ensure_packets
from side_runner import DEFAULT_CANDIDATE_COUNT, MAX_ATTEMPTS, publish_sides, run_both, run_side
from source_ingest import ingest_source
from trace_validator import validate_single_round

STOP_TASKS_PATH = Path(os.environ.get(
    "SOLOSB_STOP_TASKS_PATH",
    str(Path.home() / ".codex" / "sologsb-0917" / "stop-tasks.json"),
))


def _find_stop_marker(args: argparse.Namespace) -> str:
    data = {}
    try:
        data = json.loads(STOP_TASKS_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return ""
    markers = data.get("tasks") if isinstance(data, dict) else data
    if not isinstance(markers, list):
        return ""
    values = [str(value) for value in vars(args).values() if isinstance(value, (str, Path))]
    for raw in markers:
        marker = str(raw or "").strip()
        if marker and any(marker in value for value in values):
            return marker
    return ""


def cmd_init(args: argparse.Namespace) -> int:
    workdir = args.workdir.expanduser().resolve()
    name = args.task_name.strip() or "task"
    root = (args.task_root.expanduser().resolve() if args.task_root else workdir / safe_local_name(name))
    if root.exists() and any(root.iterdir()):
        raise SologsbError(f"任务目录已存在且不是空的: {root}")
    root.mkdir(parents=True, exist_ok=True)
    ensure_task_dirs(root)
    try:
        source_info = ingest_source(
            root / "source" / "origin",
            source=args.source,
            package=args.package,
            from_platform=args.from_platform,
            project_code=args.project_code or "",
            project_id=args.project_id or "",
            task_type=args.task_type,
            platform_base_url=args.platform_base_url or "",
            workdir=workdir,
            task_root=root,
        )
        write_json(root / "monitor" / "source.json", source_info)
        if source_info.get("mode") == "platform":
            write_json(
                root / "monitor" / "platform-selection.json",
                {
                    "status": "ready",
                    "selection": source_info.get("platformSelection") or {},
                    "platformClaim": source_info.get("platformClaim") or {},
                },
            )
        state = {
            "schemaVersion": 1,
            "status": "prepared",
            "taskName": name,
            "taskRoot": str(root),
            "createdAt": utc_now(),
            "updatedAt": utc_now(),
            "source": source_info,
            "taskType": args.task_type if args.task_type in {
                "0-1代码生成", "feature迭代", "Bug修复", "代码重构", "工程化", "代码测试"
            } else "",
            "difficulty": args.difficulty if args.difficulty in {"困难", "地狱"} else "",
        }
        save_state(root, state)
    except Exception:
        release_project_claim(root)
        raise
    print(json.dumps({"status": "prepared", "taskRoot": str(root), "source": source_info}, ensure_ascii=False, indent=2))
    return 0


def cmd_prompt(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    result = install_prompt(
        root,
        candidate_path=args.candidate.expanduser().resolve(),
        review_path=args.review.expanduser().resolve(),
        task_type=args.task_type,
        difficulty=args.difficulty,
        history_path=args.history.expanduser().resolve() if args.history else None,
        allow_over_170=args.allow_over_170,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_github_init(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    result = init_github_repo(root, repo_name=args.repo_name or "", dry_run=args.dry_run)
    if not args.dry_run:
        # Recording plans contain an absolute projectDir, so they must be
        # created only after the candidate -> A/B mapping is fixed.
        prepare_recording(root, "A")
        prepare_recording(root, "B")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_run(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    if args.base_url:
        os.environ["SOLOSB_ANTHROPIC_BASE_URL"] = str(args.base_url).strip().rstrip("/")
    if args.side == "both":
        if args.force:
            raise SologsbError("--force 只适用于 A 或 B；重跑请分别执行")
        result = run_both(
            root,
            timeout=args.timeout,
            live=args.live,
            candidate_count=args.candidates,
            attempts=args.attempts,
        )
    else:
        result = run_side(
            root,
            args.side,
            timeout=args.timeout,
            live=args.live,
            force=args.force,
            attempts=args.attempts,
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_semantic(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    sides = ("A", "B") if args.side == "both" else (args.side,)
    packets = ensure_packets(root, sides)
    print(json.dumps({"status": "ready", "sides": list(sides), "packets": packets}, ensure_ascii=False, indent=2))
    return 0


def cmd_publish(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    result = publish_sides(
        root,
        semantic_a=args.semantic_a,
        semantic_b=args.semantic_b,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_audit(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    verification = run_verification(root, args.verification_plan.expanduser().resolve() if args.verification_plan else None)
    audit = run_audit(root) if verification.get("ok") else verification
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0 if verification.get("ok") else 1


def cmd_gsb(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    result = export_gsb(
        root,
        draft_path=args.draft.expanduser().resolve(),
        review_path=args.review.expanduser().resolve(),
        accept_schema_change=args.accept_schema_change,
    )
    state = load_state(root)
    state.update(
        {
            "gsbDraftPath": str(args.draft.expanduser().resolve()),
            "gsbReviewPath": str(args.review.expanduser().resolve()),
            "gsbExportedAt": utc_now(),
        }
    )
    save_state(root, state)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_submit(args: argparse.Namespace) -> int:
    """Audit and optionally submit through the copied Python API module."""
    root = task_root_from_arg(args.task_root)
    script = Path(__file__).resolve().parent.parent / "submission" / "scripts" / "submit_api.py"
    if not script.is_file():
        raise SologsbError(f"缺少提交脚本: {script}")
    command = [sys.executable, str(script), "--task-root", str(root)]
    if args.payload:
        command += ["--payload", str(args.payload.expanduser().resolve())]
    if args.approval:
        command += ["--approval", str(args.approval.expanduser().resolve())]
    if args.execute:
        command.append("--execute")
    if args.skip_preflight:
        command.append("--skip-preflight")
    if getattr(args, "force", False):
        command.append("--force")
    command += ["--server", args.server, "--keychain-service", args.keychain_service, "--poll-timeout", str(args.poll_timeout)]
    proc = subprocess.run(command, check=False)
    return proc.returncode


def cmd_approve_line_gate(args: argparse.Namespace) -> int:
    """Create the only allowed manual exception approval (approver from device config)."""
    root = task_root_from_arg(args.task_root)
    script = Path(__file__).resolve().parent.parent / "submission" / "scripts" / "confirm_submission.py"
    if not script.is_file():
        raise SologsbError(f"缺少改动量例外审批脚本: {script}")
    command = [
        sys.executable,
        str(script),
        "--task-root",
        str(root),
        "--scope",
        "change-volume-line-gate",
    ]
    if args.payload:
        command += ["--payload", str(args.payload.expanduser().resolve())]
    if args.output:
        command += ["--output", str(args.output.expanduser().resolve())]
    return subprocess.run(command, check=False).returncode


def cmd_record(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    plan = args.plan.expanduser().resolve() if args.plan else prepare_recording(root, args.side)
    result = record_side(root, args.side, plan, lock_timeout=args.lock_timeout)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


def _remote_heads(remote_url: str) -> set[str]:
    proc = run(["git", "ls-remote", "--heads", remote_url], check=False, timeout=60, env=github_env(require_proxy=True))
    if proc.returncode != 0:
        return set()
    heads: set[str] = set()
    for line in proc.stdout.decode("utf-8", errors="replace").splitlines():
        if "\trefs/heads/" in line:
            heads.add(line.split("refs/heads/", 1)[1])
    return heads


def _excel_errors(root: Path, state: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    excel = Path(str(state.get("excelPath") or ""))
    if not excel.is_file():
        return ["缺少 Excel"]
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(excel, data_only=True, read_only=True)
        sheet = workbook["GSB提交"]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return ["Excel 缺少数据行"]
        headers = [str(value) for value in rows[0]]
        values = dict(zip(headers, rows[1]))
        schema = json.loads(SCHEMA_FALLBACK.read_text(encoding="utf-8"))
        expected_columns = len(schema.get("fields") or [])
        if len(headers) != expected_columns:
            errors.append(f"Excel 列数不是 {expected_columns}，当前 {len(headers)}")
        for key in ("A-轨迹文件", "B-轨迹文件", "A-运行录屏", "B-运行录屏"):
            value = str(values.get(key) or "")
            if not value or not Path(value).is_file():
                errors.append(f"Excel 字段 {key} 的本地文件不存在")
        if str(values.get("A-SessionID") or "") == str(values.get("B-SessionID") or ""):
            errors.append("Excel 中 A/B SessionID 相同")
        for key in ("初始环境快照", "A-产物快照", "B-产物快照"):
            value = str(values.get(key) or "")
            if not re.fullmatch(r"https://github\.com/[^/\s]+/[^/\s]+/commit/[0-9a-fA-F]{40}/?", value):
                errors.append(f"Excel 字段 {key} 不是 40 位 commit 永久链接")
        reason = str(values.get("GSB 理由") or "")
        length = len(re.sub(r"\s+", "", reason))
        if not 150 <= length <= 240:
            errors.append(f"Excel GSB 理由长度为 {length}，要求 150–240")
        if str(values.get("GSB 结论") or "") not in {"A 更好", "Same", "B 更好"}:
            errors.append("Excel GSB 结论值无效")
    except Exception as exc:
        errors.append(f"Excel 回读失败: {exc}")
    return errors


def build_status(root: Path) -> dict[str, Any]:
    state = load_state(root)
    errors: list[str] = []
    prompt_path = Path(str(state.get("promptPath") or ""))
    if not prompt_path.is_file():
        errors.append("缺少提示词")
    elif sha256_file(prompt_path) != str(state.get("promptSha256") or ""):
        errors.append("提示词哈希不匹配")
    heads: set[str] = set()
    if state.get("remoteUrl"):
        heads = _remote_heads(str(state["remoteUrl"]))
        if heads != {"main", "A", "B"}:
            errors.append(f"远端分支不是 main/A/B: {sorted(heads)}")
    for side in ("A", "B"):
        side_state = (state.get("sides") or {}).get(side) or {}
        if side_state.get("status") != "clean":
            errors.append(f"{side} 尚未干净完成")
            continue
        trace = Path(str(side_state.get("tracePath") or ""))
        if not trace.is_file():
            errors.append(f"{side} 轨迹缺失")
        elif prompt_path.is_file():
            validation = validate_single_round(
                trace,
                expected_prompt=prompt_path.read_text(encoding="utf-8"),
                expected_session_id=str(side_state.get("sessionId") or ""),
            )
            if not validation.get("ok"):
                errors.append(f"{side} 轨迹校验失败: {validation.get('errors')}")
        video = (state.get("recordings") or {}).get(side) or {}
        video_path = Path(str(video.get("videoPath") or ""))
        if not video.get("ok") or not video_path.is_file():
            errors.append(f"{side} 视频缺失")
        elif video_path.stat().st_size > 500 * 1024 * 1024:
            errors.append(f"{side} 视频超过 500MB")
        else:
            try:
                duration = float(video.get("durationSeconds") or 0)
                if duration <= 0 or duration > 90.5:
                    errors.append(f"{side} 视频时长无效或超过 90 秒: {duration}")
                width, height = video_dimensions(video_path)
                if (width, height) != (1280, 720):
                    errors.append(f"{side} 视频不是 720p，当前为 {width}x{height}")
            except Exception:
                errors.append(f"{side} 视频时长或分辨率无效")
        window_captures = video.get("windowCaptures") or []
        cursor_reports = ((video.get("cursorSuppression") or {}).get("reports") or [])
        recording_metadata = video.get("recordingMetadata") or {}
        if video.get("captureMethod") != "window-id" or not recording_isolation_ok(
            mode=str(video.get("mode") or ""),
            window_capture_reports=window_captures,
            guard_reports=cursor_reports,
            frontmost_report=recording_metadata.get("frontmostSampling"),
            service_cleanup=recording_metadata.get("serviceCleanup"),
        ):
            errors.append(f"{side} 窗口录屏隔离门禁未通过")
        if recording_metadata.get("activationPerformed") is not False or recording_metadata.get("untouched") is not True:
            errors.append(f"{side} 录屏后台/无干扰元数据门禁未通过")
        if not recording_metadata.get("focusRestores") or recording_metadata.get("focusRestoreOk") is not True:
            errors.append(f"{side} 缺少或未通过焦点恢复记录")
        if (recording_metadata.get("serviceCleanup") or {}).get("residualAppPortListeners"):
            errors.append(f"{side} 应用端口仍有残留监听进程")
        isolation_paths = [
            *(item.get("path") for item in window_captures),
            *(item.get("path") for item in cursor_reports),
            (recording_metadata.get("frontmostSampling") or {}).get("path"),
            (recording_metadata.get("serviceCleanup") or {}).get("path"),
        ]
        if any(not Path(str(path or "")).is_file() for path in isolation_paths):
            errors.append(f"{side} 窗口录屏隔离报告缺失")
        repo_value = str(side_state.get("workspacePath") or "")
        repo = Path(repo_value) if repo_value else root / "source" / side.lower()
        if repo.is_dir():
            head = run(["git", "-C", str(repo), "rev-parse", "HEAD"], check=False)
            parent = run(["git", "-C", str(repo), "rev-parse", "HEAD^"], check=False)
            if head.returncode or head.stdout.decode().strip() != str(side_state.get("artifactSnapshot") or ""):
                errors.append(f"{side} 本地产物快照与状态不一致")
            if parent.returncode or parent.stdout.decode().strip() != str(state.get("initialSnapshot") or ""):
                errors.append(f"{side} 产物父提交不是初始快照")
    side_a_state = (state.get("sides") or {}).get("A") or {}
    side_b_state = (state.get("sides") or {}).get("B") or {}
    if side_a_state.get("harnessVersion") != side_b_state.get("harnessVersion"):
        errors.append("A/B Harness 版本不一致")
    if side_a_state.get("imageDigest") != side_b_state.get("imageDigest"):
        errors.append("A/B Docker 镜像 digest 不一致")
    if side_a_state.get("declaredContextWindow") != side_b_state.get("declaredContextWindow"):
        errors.append("A/B 上下文窗口配置不一致")
    if not Path(str(state.get("evidencePath") or "")).is_file():
        errors.append("缺少审核证据")
    errors.extend(_excel_errors(root, state))
    if not Path(str(state.get("fieldGuidePath") or "")).is_file():
        errors.append("缺少字段说明")
    line_gate_path = root / "monitor" / "change-volume-line-gate.json"
    line_gate = read_json(line_gate_path, {})
    if not line_gate_path.is_file() or not line_gate.get("sides"):
        errors.append("缺少 A/B 改动量门禁记录")
    else:
        for side in ("A", "B"):
            side_gate = (line_gate.get("sides") or {}).get(side) or {}
            if side_gate.get("id") != "change-volume-line-gate" or not side_gate.get("artifactSnapshot"):
                errors.append(f"{side} 改动量门禁记录不完整")
    recordings = state.get("recordings") or {}
    if not errors and (recordings.get("A") or {}).get("ok") and (recordings.get("B") or {}).get("ok"):
        state["status"] = "complete"
        save_state(root, state)
    return {
        "status": state.get("status"),
        "ok": not errors and state.get("status") == "complete",
        "errors": list(dict.fromkeys(errors)),
        "taskRoot": str(root),
        "promptPath": str(prompt_path) if prompt_path.is_file() else "",
        "repoUrl": state.get("repoUrl", ""),
        "excelPath": state.get("excelPath", ""),
        "fieldGuidePath": state.get("fieldGuidePath", ""),
        "recordings": recordings,
        "changeVolumeLineGate": line_gate,
    }


def cmd_status(args: argparse.Namespace) -> int:
    result = build_status(task_root_from_arg(args.task_root))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


def cmd_cleanup(args: argparse.Namespace) -> int:
    root = task_root_from_arg(args.task_root)
    prefix = f"sologsb-{safe_slug(root.name)}-"
    proc = run(["docker", "ps", "-a", "--format", "{{.Names}}"], check=False)
    removed: list[str] = []
    if proc.returncode == 0:
        for name in proc.stdout.decode().splitlines():
            if name.startswith(prefix):
                run(["docker", "rm", "-f", name], check=False)
                removed.append(name)
    verify_dir = root / "monitor" / "verify"
    if verify_dir.exists() and not args.keep_verify:
        shutil.rmtree(verify_dir)
    claim_release = release_project_claim(root)
    print(
        json.dumps(
            {
                "status": "cleaned",
                "removedContainers": removed,
                "projectClaim": claim_release,
                "keptProducts": True,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    init = sub.add_parser("init", help="接入源码并创建任务目录")
    init.add_argument("--workdir", type=Path, default=Path.cwd())
    init.add_argument("--task-name", default="")
    init.add_argument("--task-root", type=Path)
    source_group = init.add_mutually_exclusive_group(required=True)
    source_group.add_argument("--source", type=Path)
    source_group.add_argument("--package", type=Path)
    source_group.add_argument("--from-platform", action="store_true")
    init.add_argument("--project-code")
    init.add_argument("--project-id")
    init.add_argument("--platform-base-url")
    init.add_argument("--task-type", default="0-1代码生成")
    init.add_argument("--difficulty", choices=["困难", "地狱"], default="困难")
    init.set_defaults(func=cmd_init)

    prompt = sub.add_parser("prompt", help="安装并校验唯一提示词")
    prompt.add_argument("--task-root", required=True)
    prompt.add_argument("--task-type", required=True)
    prompt.add_argument("--difficulty", required=True, choices=["困难", "地狱"])
    prompt.add_argument("--candidate", type=Path, required=True)
    prompt.add_argument("--review", type=Path, required=True)
    prompt.add_argument("--history", type=Path)
    prompt.add_argument("--allow-over-170", action="store_true")
    prompt.set_defaults(func=cmd_prompt)

    github = sub.add_parser("github-init", help="候选映射后创建公开 GitHub 仓库和 main/A/B")
    github.add_argument("--task-root", required=True)
    github.add_argument("--repo-name", default="")
    github.add_argument("--dry-run", action="store_true")
    github.set_defaults(func=cmd_github_init)

    run_parser = sub.add_parser("run", help="并行运行 N 个候选并将前两名映射 A/B；默认无头")
    run_parser.add_argument("--task-root", required=True)
    run_parser.add_argument("--side", required=True, choices=["A", "B", "both"])
    run_parser.add_argument("--timeout", type=float, default=7200, help="单 attempt 超时秒数，默认 7200")
    run_parser.add_argument(
        "--candidates",
        type=int,
        default=DEFAULT_CANDIDATE_COUNT,
        help="首轮并行候选数，单 Key 默认 2；前两名完成者映射为 A/B",
    )
    run_parser.add_argument(
        "--attempts",
        type=int,
        default=MAX_ATTEMPTS,
        help="每个候选的最大实际尝试次数，默认 6",
    )
    run_parser.add_argument(
        "--base-url",
        default="",
        help="Anthropic 兼容中转站 Base URL；默认取设备配置的 claude.baseUrl",
    )
    run_parser.add_argument("--force", action="store_true", help="丢弃已映射 A/B 候选的现有尝试并重跑；跨多次 force 必须手动累计实际次数")
    visibility = run_parser.add_mutually_exclusive_group()
    visibility.add_argument(
        "--live",
        action="store_true",
        dest="live",
        help="显示 A/B 实时事件；默认关闭，使用无头执行",
    )
    visibility.add_argument(
        "--no-live",
        action="store_false",
        dest="live",
        help=argparse.SUPPRESS,
    )
    run_parser.set_defaults(live=False)
    run_parser.set_defaults(func=cmd_run)

    semantic = sub.add_parser("semantic", help="生成已完成侧的语义审核包")
    semantic.add_argument("--task-root", required=True)
    semantic.add_argument("--side", choices=["A", "B", "both"], default="both")
    semantic.set_defaults(func=cmd_semantic)

    publish = sub.add_parser("publish", help="A/B 语义审核通过后原子提交并推送")
    publish.add_argument("--task-root", required=True)
    publish.add_argument("--semantic-a", type=Path, required=True)
    publish.add_argument("--semantic-b", type=Path, required=True)
    publish.set_defaults(func=cmd_publish)

    audit = sub.add_parser("audit", help="真实验证并生成证据索引")
    audit.add_argument("--task-root", required=True)
    audit.add_argument("--verification-plan", type=Path)
    audit.set_defaults(func=cmd_audit)

    gsb = sub.add_parser("gsb", help="校验 GSB 文案并导出 Excel")
    gsb.add_argument("--task-root", required=True)
    gsb.add_argument("--draft", type=Path, required=True)
    gsb.add_argument("--review", type=Path, required=True)
    gsb.add_argument("--accept-schema-change", action="store_true")
    gsb.set_defaults(func=cmd_gsb)

    submit = sub.add_parser("submit", help="审核（含 GSB 文案历史去重）并通过 Python API 提交 GSB")
    submit.add_argument("--task-root", required=True)
    submit.add_argument("--payload", type=Path)
    submit.add_argument("--approval", type=Path)
    submit.add_argument("--execute", action="store_true", help="真正上传并提交；默认只做审核和 dry-run")
    submit.add_argument("--skip-preflight", action="store_true")
    submit.add_argument("--force", action="store_true", help="忽略本地已有提交结果，允许重新提交")
    submit.add_argument("--server", default=os.environ.get("SOLO2_SERVER", "").strip())
    submit.add_argument("--keychain-service",
                        default=os.environ.get("SOLOSB_SOLO2_KEYCHAIN_SERVICE", "").strip())
    submit.add_argument("--poll-timeout", type=float, default=1800.0)
    submit.set_defaults(func=cmd_submit)

    approve_line_gate = sub.add_parser(
        "approve-line-gate",
        help="由设备配置里的审批人在 TTY 中批准唯一的改动量低于 10 行例外",
    )
    approve_line_gate.add_argument("--task-root", required=True)
    approve_line_gate.add_argument("--payload", type=Path)
    approve_line_gate.add_argument("--output", type=Path)
    approve_line_gate.set_defaults(func=cmd_approve_line_gate)

    record = sub.add_parser("record", help="执行一侧真实桌面录屏")
    record.add_argument("--task-root", required=True)
    record.add_argument("--side", required=True, choices=["A", "B"])
    record.add_argument("--plan", type=Path)
    record.add_argument(
        "--lock-timeout",
        type=float,
        default=DEFAULT_RECORDING_LOCK_TIMEOUT,
        help="等待全局录屏锁的最长秒数；默认 7200，设为 0 表示不等待",
    )
    record.set_defaults(func=cmd_record)

    status = sub.add_parser("status", help="检查最终本地交付")
    status.add_argument("--task-root", required=True)
    status.set_defaults(func=cmd_status)

    cleanup = sub.add_parser("cleanup", help="清理临时容器和验证 clone")
    cleanup.add_argument("--task-root", required=True)
    cleanup.add_argument("--keep-verify", action="store_true")
    cleanup.set_defaults(func=cmd_cleanup)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    stop_marker = _find_stop_marker(args)
    if stop_marker:
        print(
            json.dumps({"status": "stopped", "error": f"任务已在停止名单: {stop_marker}"}, ensure_ascii=False, indent=2),
            file=sys.stderr,
        )
        return 78
    try:
        return int(args.func(args))
    except SologsbError as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 2
    except subprocess.TimeoutExpired as exc:
        print(json.dumps({"status": "error", "error": f"命令超时: {exc}"}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
