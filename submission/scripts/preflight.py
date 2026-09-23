#!/usr/bin/env python3
"""Read-only preflight for the SOLO-QA GSB submission form.

This script never uploads files and never calls a write endpoint.
"""
from __future__ import annotations

import argparse
import contextlib
import difflib
import hashlib
import importlib
import json
import os
import platform
import re
import shutil
import socket
import subprocess
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

# --- 设备配置注入：必须在下列模块读取环境变量之前执行 ---
for _parent in Path(__file__).resolve().parents:
    if (_parent / "scripts" / "device_config.py").is_file():
        sys.path.insert(0, str(_parent / "scripts"))
        break
import device_config as _device_config  # noqa: E402

_device_config.load_and_apply()

try:
    from openpyxl import load_workbook
except Exception as exc:  # pragma: no cover - dependency diagnostics
    load_workbook = None
    OPENPYXL_ERROR = str(exc)
else:
    OPENPYXL_ERROR = ""


TRACE_MAX = 27 * 1024 * 1024
VIDEO_MAX = 500 * 1024 * 1024
VIDEO_EXTS = {".mp4", ".mov", ".webm", ".m4v"}
HARNESS_VERSION = "2.1.197"
def gsb_server() -> str:
    """SOLO2 平台地址。技能包里不留默认域名，必须由设备配置提供。"""
    value = os.environ.get("SOLO2_SERVER", "").strip().rstrip("/")
    if not value:
        raise RuntimeError(
            "缺少 SOLO2 平台地址：请在设备配置里设置 solo2.baseUrl，"
            "或运行 scripts/configure.py wizard"
        )
    return value

GSB_HISTORY_CACHE_ENV = "SOLOGBS_0917_GSB_HISTORY_CACHE"
GSB_HISTORY_MANUAL_CACHE_ENV = "SOLOGBS_0917_GSB_HISTORY_MANUAL_CACHE"
GSB_HISTORY_CACHE_TTL = float(os.environ.get("SOLOGBS_0917_GSB_HISTORY_CACHE_TTL", "300") or "300")
REASON_FRAGMENT_MIN = int(os.environ.get("SOLOGBS_REASON_FRAGMENT_MIN", "12") or "12")
REASON_FRAGMENT_REVIEW = int(os.environ.get("SOLOGBS_REASON_FRAGMENT_REVIEW", "8") or "8")
REASON_SIMILARITY_BLOCK = float(os.environ.get("SOLOGBS_REASON_SIMILARITY_BLOCK", "0.40") or "0.40")
REASON_SIMILARITY_REVIEW = float(os.environ.get("SOLOGBS_REASON_SIMILARITY_REVIEW", "0.20") or "0.20")
REASON_NGRAM_SIZE = int(os.environ.get("SOLOGBS_REASON_NGRAM_SIZE", "6") or "6")
REASON_NGRAM_BLOCK = float(os.environ.get("SOLOGBS_REASON_NGRAM_BLOCK", "0.45") or "0.45")
REASON_NGRAM_REVIEW = float(os.environ.get("SOLOGBS_REASON_NGRAM_REVIEW", "0.20") or "0.20")
COMMIT_RE = re.compile(r"^https://github\.com/[^/\s]+/[^/\s]+/commit/[0-9a-fA-F]{40}/?$")
OBJECTIVE_EVIDENCE_RE = re.compile(
    r"(?i)(?:[A-Za-z0-9_.-]+\.(?:go|ts|tsx|vue|py|js|json|ya?ml|md|sql|sh|css|html)|"
    r"`[^`]+`|(?:Exit code|HTTP|返回码|退出码|报错|错误|Exception|Error)\s*[:：]?\s*\S+)"
)
OBJECTIVE_CONSEQUENCE_RE = re.compile(
    r"(导致|造成|没有|未|失败|报错|退出码|无法|阻断|拒绝|缺少|缺失|未落地|未启动|未接入|未完成|影响)"
)
MACHINE_STEP_RE = re.compile(r"第\s*\d{3,}\s*次")
PROCESS_ACTION_RE = re.compile(
    r"(?:阅读|读取|查阅|浏览|查看|检查|核对|定位|搜索|检索|梳理|比对|修改|编辑|调整|重构|"
    r"补充|补齐|补全|同步|运行|执行|调用|调试|排查|验证|测试|构建|启动|安装|新增|删除|实现|发现|修复)"
)
PROCESS_LOCATOR_RE = re.compile(
    r"(?:第\s*[一二三四五六七八九十\d]+\s*(?:步|次|轮|阶段)|阶段|初始化|联调|验证|构建|启动|"
    r"迁移|数据库|脚本|文件|接口|服务|模块|路由|中间件|指标|功能|字段|页面|流程|逻辑|配置|模型|需求|用例|`[^`]+`|"
    r"\.(?:go|js|cjs|mjs|ts|tsx|jsx|py|java|kt|rs|vue|json|ya?ml|toml|md|sql|sh|css|html|xml))"
)
ARTIFACT_OUTCOME_RE = re.compile(
    r"(?:返回|输出|缺少|缺失|未实现|没有|失败|报错|异常|保留|仍然|仍含|写入|生成|创建|删除|更新|"
    r"展示|完成|支持|正常|通过|可用|实现|拒绝|500|404)"
)
FIELD_FACTOR_RE = re.compile(
    r"(?:录屏|屏幕录制|录制过程|录制画面|录制结果|视频画面|视频中|视频里|视频显示|视频可见|视频证据|"
    r"视频|截图|画面中|画面显示|画面可见|镜头|剪辑|剪掉|\bOtty\b|\biTerm2?\b|"
    r"1280\s*[x×]\s*720|720p|\bMP4\b|鼠标|光标|终端窗口|终端界面|命令行窗口|"
    r"浏览器|浏览器窗口|屏幕|测试设备|测试机|运行环境|运行机器|验收宿主|验收机|采集环境|采集设备|录制设备)",
    re.I,
)
NON_CONTAINER_TEST_ARTIFACT_RE = re.compile(
    r"(?:本地|容器外|非容器内|宿主(?:机|侧)?)[^。；，]{0,16}"
    r"(?:测试文件|测试脚本|测试代码|自测文件|自测脚本|验收脚本|验收测试文件|验收测试脚本)|"
    r"(?:验收脚本|验收测试文件|验收测试脚本)|"
    r"(?:测试文件|测试脚本|测试代码|自测文件|自测脚本)[^。；，]{0,16}"
    r"(?:由本地|在容器外|非容器内|不是容器内|容器外生成|本地生成)",
    re.I,
)
CHANGE_VOLUME_APPROVAL_SCOPE = "change-volume-line-gate"
AUTO_APPROVER = os.environ.get("SOLOGBS_AUTO_APPROVER", "").strip() or "auto"
BANNED_REASON_PATTERN = re.compile(r"落在.{0,16}")
# 文案通则里的禁用词、排版与文风规则统一委托给 scripts/gsb_tools.py，避免两处口径漂移。
SKILL_SCRIPTS_DIR = Path(__file__).resolve().parents[2] / "scripts"


def _load_reason_style_helpers():
    try:
        if str(SKILL_SCRIPTS_DIR) not in sys.path:
            sys.path.insert(0, str(SKILL_SCRIPTS_DIR))
        return importlib.import_module("gsb_tools")
    except Exception:
        return None

CODE_VOLUME_HARD_MIN = 10
CODE_VOLUME_TARGET = 30
SOURCE_EXTENSIONS = {
    ".c", ".cc", ".cpp", ".cs", ".css", ".go", ".graphql", ".h", ".hpp", ".html",
    ".java", ".js", ".jsx", ".kt", ".kts", ".less", ".m", ".mm", ".mjs", ".cjs",
    ".php", ".proto", ".py", ".rb", ".rs", ".scala", ".scss", ".sh", ".sql",
    ".svelte", ".swift", ".ts", ".tsx", ".vue",
}
GENERATED_DIR_NAMES = {
    "node_modules", "dist", "build", "coverage", ".next", ".nuxt", ".vite",
    "target", "vendor", "__pycache__", ".venv", "venv", ".cache",
}


def _proxy_port_open(host: str, port: int, timeout: float = 0.25) -> bool:
    try:
        with socket.create_connection((host, port), timeout=timeout):
            return True
    except OSError:
        return False


def github_env(extra: dict[str, str] | None = None, *, require_proxy: bool = False) -> dict[str, str]:
    """Apply the local Loon proxy to GitHub CLI and Git network commands."""
    env = os.environ.copy()
    proxy = (
        os.environ.get("SOLOSB_GITHUB_PROXY", "").strip()
        or os.environ.get("GITHUB_PROXY", "").strip()
    )
    if not proxy:
        if _proxy_port_open("127.0.0.1", 17890):
            proxy = "http://127.0.0.1:17890"
        elif _proxy_port_open("127.0.0.1", 17891):
            proxy = "socks5h://127.0.0.1:17891"
    if not proxy and require_proxy:
        raise RuntimeError(
            "Loon GitHub 代理不可用：请设置 SOLOSB_GITHUB_PROXY，"
            "或启动 HTTP 127.0.0.1:17890 / SOCKS5 127.0.0.1:17891"
        )
    if proxy:
        env["HTTPS_PROXY"] = proxy
        env["HTTP_PROXY"] = proxy
        env["ALL_PROXY"] = proxy
        no_proxy = env.get("NO_PROXY", "")
        required = ["127.0.0.1", "localhost", "::1"]
        merged = [item for item in no_proxy.split(",") if item.strip()]
        for item in required:
            if item not in merged:
                merged.append(item)
        env["NO_PROXY"] = ",".join(merged)
        env["no_proxy"] = env["NO_PROXY"]
    if extra:
        env.update({str(key): str(value) for key, value in extra.items()})
    return env
LOCKFILE_NAMES = {
    "package-lock.json", "pnpm-lock.yaml", "yarn.lock", "bun.lockb",
    "composer.lock", "Gemfile.lock", "poetry.lock", "Pipfile.lock",
    "Cargo.lock", "go.sum",
}
DOC_EXTENSIONS = {".md", ".rst", ".txt", ".adoc"}


def _is_generated_or_lock_path(path: str) -> bool:
    parts = [part for part in Path(path).parts if part not in {"", "."}]
    name = parts[-1] if parts else ""
    return any(part in GENERATED_DIR_NAMES for part in parts) or name in LOCKFILE_NAMES


def _is_source_path(path: str) -> bool:
    return Path(path).suffix.lower() in SOURCE_EXTENSIONS


def _is_test_path(path: str) -> bool:
    parts = [part.lower() for part in Path(path).parts]
    name = parts[-1] if parts else ""
    return (
        any(part in {"test", "tests", "__tests__", "testdata", "spec"} for part in parts)
        or ".test." in name
        or ".spec." in name
        or name.endswith("_test.go")
        or name.startswith("test_")
    )


def _filtered_numstat(repo: Path, base: str, head: str, paths: list[str]) -> tuple[int, dict[str, int]]:
    total = 0
    per_file: dict[str, int] = {}
    for offset in range(0, len(paths), 80):
        chunk = paths[offset : offset + 80]
        proc = run(
            ["git", "-C", str(repo), "diff", "--numstat", base, head, "--", *chunk],
            timeout=300,
        )
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.strip() or f"git diff --numstat 失败: {base}..{head}")
        for line in proc.stdout.splitlines():
            parts = line.split("\t", 2)
            if len(parts) != 3 or parts[0] == "-" or parts[1] == "-":
                continue
            added, deleted, path = int(parts[0]), int(parts[1]), parts[2]
            count = max(added, deleted)
            total += count
            per_file[path] = per_file.get(path, 0) + count
    return total, per_file


def verify_code_change(state: dict, remote: dict, task_root: Path | None = None) -> dict:
    """Approximate platform G11 locally against the published A/B commits."""
    task_root = Path(task_root or state.get("taskRoot") or ".").expanduser().resolve()
    result = {
        "ok": False,
        "hardMinimumLines": CODE_VOLUME_HARD_MIN,
        "targetLines": CODE_VOLUME_TARGET,
        "targetFiles": 3,
        "sides": {},
        "badPaths": [],
        "errors": [],
    }
    remote_url = str(state.get("remoteUrl") or "")
    initial = str(state.get("initialSnapshot") or "").lower()
    heads = {str(key): str(value).lower() for key, value in (remote.get("heads") or {}).items()}
    main_branch = "main" if "main" in heads else ("master" if "master" in heads else "")
    if not remote_url or not initial or not main_branch:
        result["errors"].append("远端 URL、初始快照或主分支缺失")
        return result
    for side in ("A", "B"):
        if not heads.get(side):
            result["errors"].append(f"{side} 分支 SHA 缺失")
    if result["errors"]:
        return result

    import tempfile

    local_repo = task_root / "source" / "origin"
    use_local_repo = False
    if local_repo.is_dir():
        expected_refs = {
            main_branch: heads[main_branch],
            "A": heads["A"],
            "B": heads["B"],
        }
        local_refs: dict[str, str] = {}
        for branch, expected in expected_refs.items():
            ref = "refs/remotes/origin/" + branch
            proc = run(["git", "-C", str(local_repo), "rev-parse", "--verify", ref], timeout=20)
            if proc.returncode == 0:
                local_refs[branch] = str(proc.stdout).strip().lower()
        use_local_repo = local_refs == {key: str(value).lower() for key, value in expected_refs.items()}

    manager = contextlib.nullcontext(local_repo) if use_local_repo else tempfile.TemporaryDirectory(prefix="gsb-code-volume-")
    with manager as temp:
        repo = local_repo if use_local_repo else Path(temp) / "repo"
        if not use_local_repo:
            clone = run(
                [
                    "git", "clone", "--filter=blob:none", "--no-checkout",
                    "--single-branch", "--branch", main_branch, remote_url, str(repo),
                ],
                timeout=600,
                env=github_env(require_proxy=True),
            )
            if clone.returncode != 0:
                result["errors"].append(clone.stderr.strip() or "代码改动量检查 clone 失败")
                return result
            fetch = run(
                [
                    "git", "-C", str(repo), "fetch", "--filter=blob:none", "origin",
                    f"{heads['A']}:refs/remotes/origin/A",
                    f"{heads['B']}:refs/remotes/origin/B",
                ],
                timeout=600,
                env=github_env(require_proxy=True),
            )
            if fetch.returncode != 0:
                result["errors"].append(fetch.stderr.strip() or "代码改动量检查 fetch A/B 失败")
                return result

        bad_paths: set[str] = set()
        for side in ("A", "B"):
            head = heads[side]
            names = run(
                ["git", "-C", str(repo), "diff", "--name-only", "--diff-filter=ACMRTUXB", initial, head],
                timeout=300,
            )
            if names.returncode != 0:
                result["errors"].append(names.stderr.strip() or f"{side} 变更文件枚举失败")
                continue
            changed = [line.strip() for line in names.stdout.splitlines() if line.strip()]
            bad = [path for path in changed if _is_generated_or_lock_path(path)]
            bad_paths.update(bad)
            business = [
                path for path in changed
                if _is_source_path(path)
                and not _is_test_path(path)
                and not _is_generated_or_lock_path(path)
                and Path(path).suffix.lower() not in DOC_EXTENSIONS
            ]
            try:
                lines, per_file = _filtered_numstat(repo, initial, head, business)
            except Exception as exc:
                result["errors"].append(f"{side}: {exc}")
                continue
            top_files = [
                {"path": path, "lines": count}
                for path, count in sorted(per_file.items(), key=lambda item: (-item[1], item[0]))[:20]
            ]
            result["sides"][side] = {
                "head": head,
                "changedFileCount": len(changed),
                "businessFiles": business,
                "businessFileCount": len(business),
                "lines": lines,
                "topFiles": top_files,
                "hardOk": lines >= CODE_VOLUME_HARD_MIN,
                "targetOk": lines >= CODE_VOLUME_TARGET and len(business) >= 3,
                "generatedOrLockPaths": bad[:100],
            }
        result["badPaths"] = sorted(bad_paths)
        result["okHygiene"] = not result["badPaths"]
        result["ok"] = bool(result["sides"]) and all(item.get("hardOk") for item in result["sides"].values()) and not result["errors"]
        return result


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_json(value: object) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def run(
    args: list[str],
    *,
    timeout: int = 60,
    env: dict[str, str] | None = None,
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        args,
        capture_output=True,
        text=True,
        check=False,
        timeout=timeout,
        env=env,
    )


def find_task_root(explicit: Path | None) -> Path:
    if explicit:
        return explicit.expanduser().resolve()
    current = Path.cwd().resolve()
    for candidate in (current, *current.parents):
        if (candidate / "monitor" / "state.json").is_file():
            return candidate
    raise SystemExit("未找到 monitor/state.json；请传入 --task-root")


def load_json(path: Path) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def page_schema() -> dict:
    path = Path(__file__).resolve().parents[1] / "references" / "gsb-submit-schema.json"
    return load_json(path)

def _keychain_secret(service: str) -> str:
    account = os.environ.get("USER", "")
    proc = subprocess.run(
        ["security", "find-generic-password", "-a", account, "-s", service, "-w"],
        capture_output=True,
        text=True,
        check=False,
    )
    return proc.stdout.strip() if proc.returncode == 0 else ""


def _readonly_json(path: str, *, timeout: int = 60) -> dict:
    def _call(cookie: str, csrf: str) -> dict:
        request = urllib.request.Request(
            gsb_server() + path,
            headers={
                "Accept": "application/json",
                "Cookie": cookie,
                "x-csrf-token": csrf,
                "Referer": gsb_server() + "/app/gsb/submissions",
                "User-Agent": "gsb-submit-preflight/1.0",
            },
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8") or "{}")

    service = os.environ.get("SOLOSB_SOLO2_KEYCHAIN_SERVICE", "").strip()
    cookie = os.environ.get("SOLO_QA_COOKIE", "").strip() or (
        _keychain_secret(service + "-cookie") if service else ""
    )
    csrf = os.environ.get("SOLO_QA_CSRF", "").strip() or (
        _keychain_secret(service + "-csrf") if service else ""
    )
    if not cookie or not csrf:
        if _device_config.refresh_solo2_into_env():
            cookie = os.environ.get("SOLO_QA_COOKIE", "").strip()
            csrf = os.environ.get("SOLO_QA_CSRF", "").strip()
    if not cookie or not csrf:
        raise RuntimeError("缺少 SOLO-QA 只读凭据")
    try:
        return _call(cookie, csrf)
    except urllib.error.HTTPError as exc:
        if exc.code == 401 and _device_config.refresh_solo2_into_env():
            try:
                return _call(os.environ["SOLO_QA_COOKIE"], os.environ["SOLO_QA_CSRF"])
            except urllib.error.HTTPError as retry_exc:
                body = retry_exc.read().decode("utf-8", errors="replace")
                raise RuntimeError(
                    f"重新登录后仍返回 HTTP {retry_exc.code}: {body[:1000]}"
                ) from retry_exc
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"只读请求 HTTP {exc.code}: {body[:1000]}") from exc


def gsb_history_cache_path() -> Path:
    configured = os.environ.get(GSB_HISTORY_CACHE_ENV, "").strip()
    if configured:
        return Path(configured).expanduser().resolve()
    codex_home = Path(os.environ.get("CODEX_HOME") or (Path.home() / ".codex")).expanduser()
    return (codex_home / "cache" / "sologsb-0917" / "gsb-history-cache.json").resolve()


def gsb_history_manual_cache_path() -> Path:
    configured = os.environ.get(GSB_HISTORY_MANUAL_CACHE_ENV, "").strip()
    if configured:
        return Path(configured).expanduser().resolve()
    codex_home = Path(os.environ.get("CODEX_HOME") or (Path.home() / ".codex")).expanduser()
    return (codex_home / "cache" / "sologsb-0917" / "gsb-history-manual.json").resolve()


def _normalize_history_item(item: dict) -> dict | None:
    if not isinstance(item, dict):
        return None
    raw_id = item.get("id") or item.get("submissionNo") or item.get("submission_id")
    try:
        item_id = int(raw_id or 0)
    except (TypeError, ValueError):
        item_id = 0
    prompt = str(item.get("userPrompt") or item.get("user_prompt") or "").strip()
    reason = str(item.get("gsbReason") or item.get("gsb_reason") or item.get("reason") or "").strip()
    if item_id <= 0:
        return None
    return {
        "id": item_id,
        "status": str(item.get("status") or "manual"),
        "submittedAt": str(item.get("submittedAt") or item.get("submitted_at") or ""),
        "userPrompt": prompt,
        "gsbReason": reason,
        "questionType": str(item.get("questionType") or item.get("question_type") or ""),
        "difficulty": str(item.get("difficulty") or ""),
        "languages": str(item.get("languages") or ""),
        "repoId": str(item.get("repoId") or item.get("repo_id") or ""),
        "aSessionId": str(item.get("aSessionId") or item.get("a_session_id") or ""),
        "bSessionId": str(item.get("bSessionId") or item.get("b_session_id") or ""),
        "aDescDelivery": str(item.get("aDescDelivery") or item.get("a_desc_delivery") or ""),
        "bDescDelivery": str(item.get("bDescDelivery") or item.get("b_desc_delivery") or ""),
        "source": "manual",
        "unresolved": not prompt and not reason,
    }


def load_manual_history_items() -> list[dict]:
    path = gsb_history_manual_cache_path()
    raw = load_json(path)
    values = raw.get("items") if isinstance(raw, dict) else raw
    if isinstance(values, dict):
        values = [values]
    if not isinstance(values, list):
        return []
    items = []
    for value in values:
        normalized = _normalize_history_item(value)
        if normalized:
            items.append(normalized)
    return items


def merge_manual_history_items(history: dict) -> dict:
    manual = load_manual_history_items()
    result = dict(history)
    by_id = {
        int(item.get("id") or 0): dict(item)
        for item in result.get("items") or []
        if int(item.get("id") or 0) > 0
    }
    for item in manual:
        by_id[int(item["id"])] = item
    result["items"] = sorted(by_id.values(), key=lambda item: int(item.get("id") or 0), reverse=True)
    result["total"] = len(result["items"])
    result["manualCachePath"] = str(gsb_history_manual_cache_path())
    result["manualCount"] = len(manual)
    return result


def write_manual_history_items(items: list[dict]) -> Path:
    normalized = []
    for item in items:
        value = _normalize_history_item(item)
        if value:
            normalized.append(value)
    path = gsb_history_manual_cache_path()
    _write_json_atomic(path, {"source": "platform-qc-or-audit-feedback", "updatedAt": utc_now(), "items": normalized})
    return path


def _write_json_atomic(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temp, path)


def _cache_age_seconds(value: dict) -> float | None:
    fetched_at = str(value.get("fetchedAt") or "").strip()
    if not fetched_at:
        return None
    try:
        parsed = datetime.fromisoformat(fetched_at.replace("Z", "+00:00"))
    except ValueError:
        return None
    return max(0.0, (datetime.now(timezone.utc) - parsed).total_seconds())


def _filter_gsb_history(
    history: dict,
    current_sessions: set[str],
    exclude_ids: set[int],
) -> dict:
    items = []
    skipped_current = []
    for item in history.get("items") or []:
        if not isinstance(item, dict):
            continue
        item_id = int(item.get("id") or 0)
        sessions = {str(item.get("aSessionId") or ""), str(item.get("bSessionId") or "")} - {""}
        if item_id in exclude_ids or (sessions and sessions <= current_sessions):
            skipped_current.append(item_id)
            continue
        if (
            not str(item.get("userPrompt") or "").strip()
            and not str(item.get("gsbReason") or "").strip()
            and not item.get("unresolved")
        ):
            continue
        items.append(dict(item))
    result = dict(history)
    result["total"] = len(items)
    result["items"] = items
    result["skippedCurrentSubmissions"] = sorted(set(skipped_current))
    return result


def _fetch_live_gsb_history() -> dict:
    params = {
        "verdict": "", "git_state": "", "difficulty": "", "needs_review": "false",
        "keyword": "", "date_from": "", "date_to": "", "user_id": "0", "team": "",
        "ids": "", "leader_id": "0", "question_type": "", "only_stale": "false",
        "stage": "", "page": "1", "page_size": "200",
    }
    first = _readonly_json("/api/v1/gsb/submissions?" + urllib.parse.urlencode(params))
    meta = first.get("meta") or {}
    total_pages = int(meta.get("total_pages") or 1)
    list_items = list(first.get("items") or [])
    for page in range(2, total_pages + 1):
        params["page"] = str(page)
        payload = _readonly_json("/api/v1/gsb/submissions?" + urllib.parse.urlencode(params))
        list_items.extend(payload.get("items") or [])

    items = []
    for item in list_items:
        item_id = int(item.get("id") or 0)
        if item_id <= 0:
            continue
        detail = _readonly_json(f"/api/v1/gsb/submissions/{item_id}")
        prompt = str(detail.get("user_prompt") or "").strip()
        reason = str(detail.get("gsb_reason") or detail.get("reason") or "").strip()
        if not prompt and not reason:
            continue
        items.append({
            "id": item_id,
            "status": str(detail.get("status_label") or detail.get("status") or ""),
            "submittedAt": str(detail.get("submitted_at") or detail.get("created_at") or ""),
            "userPrompt": prompt,
            "gsbReason": reason,
            "questionType": str(detail.get("question_type") or ""),
            "difficulty": str(detail.get("difficulty") or ""),
            "languages": str(detail.get("languages") or ""),
            "repoId": str(detail.get("repo_id") or ""),
            "aSessionId": str(detail.get("a_session_id") or ""),
            "bSessionId": str(detail.get("b_session_id") or ""),
            "aDescDelivery": str(detail.get("a_desc_delivery") or "").strip(),
            "bDescDelivery": str(detail.get("b_desc_delivery") or "").strip(),
        })
    return {
        "source": gsb_server() + "/app/gsb/submissions",
        "fetchedAt": utc_now(),
        "serverTotal": int(meta.get("total") or len(list_items)),
        "items": items,
    }


def unresolved_history_ids(history: dict) -> list[int]:
    return sorted({
        int(item.get("id") or 0)
        for item in history.get("items") or []
        if isinstance(item, dict) and item.get("unresolved") and int(item.get("id") or 0) > 0
    })


def fetch_gsb_prompt_history(
    current_sessions: set[str],
    exclude_ids: set[int],
    *,
    force_refresh: bool = False,
    max_age_seconds: float = GSB_HISTORY_CACHE_TTL,
) -> dict:
    cache_path = gsb_history_cache_path()
    cache = load_json(cache_path)
    age = _cache_age_seconds(cache)
    if cache.get("items") and not force_refresh and age is not None and age <= max_age_seconds:
        result = _filter_gsb_history(merge_manual_history_items(cache), current_sessions, exclude_ids)
        result.update({
            "cachePath": str(cache_path),
            "cacheHit": True,
            "cacheFresh": True,
            "cacheAgeSeconds": round(age, 3),
            "fetchError": "",
        })
        return result
    try:
        live = _fetch_live_gsb_history()
        _write_json_atomic(cache_path, live)
        result = _filter_gsb_history(merge_manual_history_items(live), current_sessions, exclude_ids)
        result.update({
            "cachePath": str(cache_path),
            "cacheHit": False,
            "cacheFresh": True,
            "cacheAgeSeconds": 0.0,
            "fetchError": "",
        })
        return result
    except Exception as exc:
        manual_count = len(load_manual_history_items())
        if not cache.get("items") and not manual_count:
            raise
        result = _filter_gsb_history(merge_manual_history_items(cache), current_sessions, exclude_ids)
        result.update({
            "cachePath": str(cache_path),
            "cacheHit": True,
            "cacheFresh": False,
            "cacheAgeSeconds": round(age, 3) if age is not None else None,
            "fetchError": str(exc),
        })
        return result


# Backward-compatible alias used by the standalone prompt_dedup CLI.
fetch_gsb_history = fetch_gsb_prompt_history


def _normalize_prompt(value: str) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"\s+", "", text).casefold()
    return re.sub(r"[，。；：！？、,.!?;:\"'“”‘’（）()《》<>【】\[\]{}…—_-]", "", text)


def _longest_common_substring(left: str, right: str) -> int:
    if not left or not right:
        return 0
    previous = [0] * (len(right) + 1)
    best = 0
    for lchar in left:
        current = [0]
        for index, rchar in enumerate(right, 1):
            value = previous[index - 1] + 1 if lchar == rchar else 0
            current.append(value)
            best = max(best, value)
        previous = current
    return best


def _ngram_overlap(left: str, right: str, size: int = REASON_NGRAM_SIZE) -> tuple[float, float]:
    if not left or not right:
        return 0.0, 0.0
    left_set = {left[index:index + size] for index in range(max(0, len(left) - size + 1))}
    right_set = {right[index:index + size] for index in range(max(0, len(right) - size + 1))}
    if not left_set or not right_set:
        return 0.0, 0.0
    intersection = len(left_set & right_set)
    return intersection / min(len(left_set), len(right_set)), intersection / len(left_set | right_set)


def assess_prompt_dedup(candidate: str, history: dict) -> dict:
    normalized = _normalize_prompt(candidate)
    matches = []
    for item in history.get("items") or []:
        other = _normalize_prompt(item.get("userPrompt") or "")
        similarity = difflib.SequenceMatcher(None, normalized, other).ratio() if normalized or other else 1.0
        common = _longest_common_substring(normalized, other)
        exact = bool(normalized) and normalized == other
        decision = "UNIQUE"
        if exact:
            decision = "EXACT"
        elif common >= 18 or similarity >= 0.82:
            decision = "SIMILAR"
        elif similarity >= 0.62:
            decision = "REVIEW_REQUIRED"
        if decision != "UNIQUE":
            matches.append({
                "id": item.get("id"),
                "decision": decision,
                "similarity": round(similarity, 4),
                "longestCommonSubstringLength": common,
                "userPrompt": item.get("userPrompt"),
                "questionType": item.get("questionType"),
                "difficulty": item.get("difficulty"),
                "repoId": item.get("repoId"),
                "status": item.get("status"),
            })
    matches.sort(key=lambda item: (item["decision"] != "EXACT", -item["similarity"], -item["longestCommonSubstringLength"]))
    if any(item["decision"] == "EXACT" for item in matches):
        decision = "EXACT"
    elif any(item["decision"] == "SIMILAR" for item in matches):
        decision = "SIMILAR"
    elif matches:
        decision = "REVIEW_REQUIRED"
    else:
        decision = "UNIQUE"
    return {
        "decision": decision,
        "duplicate": decision in {"EXACT", "SIMILAR"},
        "reviewRequired": decision == "REVIEW_REQUIRED",
        "candidateSha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "historyCount": len(history.get("items") or []),
        "matches": matches[:20],
    }


def assess_gsb_reason_dedup(candidate: str, history: dict) -> dict:
    normalized = _normalize_prompt(candidate)
    if not normalized:
        return {
            "decision": "MISSING",
            "duplicate": True,
            "reviewRequired": True,
            "candidateSha256": "",
            "historyCount": len(history.get("items") or []),
            "matches": [],
            "rewriteInstruction": "GSB 理由为空，先依据实际轨迹与产物重写。",
        }
    matches = []
    for item in history.get("items") or []:
        other = _normalize_prompt(item.get("gsbReason") or "")
        if not other:
            continue
        similarity = difflib.SequenceMatcher(None, normalized, other).ratio()
        common = _longest_common_substring(normalized, other)
        ngram_containment, ngram_jaccard = _ngram_overlap(normalized, other)
        exact = normalized == other
        decision = "UNIQUE"
        rule = ""
        if exact:
            decision = "EXACT"
            rule = "EXACT"
        elif common >= REASON_FRAGMENT_MIN or similarity >= REASON_SIMILARITY_BLOCK or ngram_containment >= REASON_NGRAM_BLOCK:
            decision = "SIMILAR"
            if common >= REASON_FRAGMENT_MIN:
                rule = "B-5.long_fragment"
            elif ngram_containment >= REASON_NGRAM_BLOCK:
                rule = "B-5.template_ngram"
            else:
                rule = "B-5.high_similarity"
        elif (
            common >= REASON_FRAGMENT_REVIEW
            or ngram_containment >= REASON_NGRAM_REVIEW
            or (similarity >= REASON_SIMILARITY_REVIEW and common >= 4)
        ):
            decision = "REVIEW_REQUIRED"
            if common >= REASON_FRAGMENT_REVIEW:
                rule = "B-5.long_fragment"
            elif ngram_containment >= REASON_NGRAM_REVIEW:
                rule = "B-5.template_ngram"
            else:
                rule = "B-5.low_similarity"
        if decision != "UNIQUE":
            matches.append({
                "id": item.get("id"),
                "decision": decision,
                "rule": rule,
                "similarity": round(similarity, 4),
                "similarityPercent": round(similarity * 100, 1),
                "longestCommonSubstringLength": common,
                "ngramContainment": round(ngram_containment, 4),
                "ngramJaccard": round(ngram_jaccard, 4),
                "gsbReason": item.get("gsbReason"),
                "questionType": item.get("questionType"),
                "difficulty": item.get("difficulty"),
                "repoId": item.get("repoId"),
                "status": item.get("status"),
            })
    matches.sort(key=lambda item: (item["decision"] != "EXACT", -item["similarity"], -item["longestCommonSubstringLength"]))
    if any(item["decision"] == "EXACT" for item in matches):
        decision = "EXACT"
    elif any(item["decision"] == "SIMILAR" for item in matches):
        decision = "SIMILAR"
    elif matches:
        decision = "REVIEW_REQUIRED"
    else:
        decision = "UNIQUE"
    instruction = ""
    if matches:
        top = matches[0]
        instruction = (
            f"GSB 理由与历史 #{top.get('id')} 相似度 {top.get('similarityPercent')}%，"
            f"命中 {top.get('rule') or 'B-5'}；必须针对本次与历史记录两次运行的实际轨迹、commit、"
            "测试或启动输出重写对比理由，不得沿用公共句式或只替换项目名。"
        )
    return {
        "decision": decision,
        "duplicate": decision != "UNIQUE",
        "reviewRequired": decision != "UNIQUE",
        "candidateSha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "historyCount": len(history.get("items") or []),
        "rewriteInstruction": instruction,
        "matches": matches[:20],
    }

def assess_delivery_dedup(descriptions: dict[str, str], history: dict) -> dict:
    """G12：两段交付完整性描述分别与历史 A/B 两侧描述比对。

    描述通常较短，公共短语多，REVIEW_REQUIRED 只作警告；EXACT/SIMILAR 阻断。
    """
    pool = []
    for item in history.get("items") or []:
        for key in ("aDescDelivery", "bDescDelivery"):
            text = str(item.get(key) or "").strip()
            if text:
                pool.append({**item, "gsbReason": text, "matchedField": key})
    sides = {}
    for side, text in descriptions.items():
        result = assess_gsb_reason_dedup(text, {"items": pool})
        result["rewriteInstruction"] = (result.get("rewriteInstruction") or "").replace("GSB 理由", f"{side}-交付完整性描述")
        sides[side] = result
    decisions = [value.get("decision") for value in sides.values()]
    if "MISSING" in decisions:
        decision = "MISSING"
    elif "EXACT" in decisions:
        decision = "EXACT"
    elif "SIMILAR" in decisions:
        decision = "SIMILAR"
    elif "REVIEW_REQUIRED" in decisions:
        decision = "REVIEW_REQUIRED"
    else:
        decision = "UNIQUE"
    return {"decision": decision, "historyCount": len(pool), "sides": sides}


def expected_os() -> str:
    system = platform.system().lower()
    if system == "windows":
        return "Windows"
    return "MacOS/Linux"


def parse_commit_sha(url: str) -> str:
    match = re.search(r"/commit/([0-9a-fA-F]{40})/?$", str(url or ""))
    return match.group(1).lower() if match else ""


def parse_remote(url: str) -> tuple[str, str]:
    value = str(url or "").strip()
    if value.startswith("git@github.com:"):
        value = "https://github.com/" + value.split(":", 1)[1]
    parsed = urlparse(value)
    parts = [part for part in parsed.path.strip("/").split("/") if part]
    if len(parts) >= 2:
        return parts[0], parts[1].removesuffix(".git")
    return "", ""


def derived_expected(state: dict, draft: dict) -> dict[str, str]:
    sides = state.get("sides") or {}
    a = sides.get("A") or {}
    b = sides.get("B") or {}
    recordings = state.get("recordings") or {}
    delivery = draft.get("delivery") if isinstance(draft.get("delivery"), dict) else {}
    delivery_a = delivery.get("A") or {}
    delivery_b = delivery.get("B") or {}
    return {
        "user_prompt": str(state.get("promptText") or "").strip(),
        "question_type": str(state.get("taskType") or ""),
        "difficulty": str(state.get("difficulty") or ""),
        "languages": str(draft.get("languages") or ""),
        "harness": str(a.get("harness") or b.get("harness") or ""),
        "harness_version": HARNESS_VERSION,
        "os_platform": expected_os(),
        "repro_level": str(draft.get("repro_level") or ""),
        "env_snapshot": str(state.get("initialSnapshotUrl") or ""),
        "a_session_id": str(a.get("sessionId") or ""),
        "a_trace_file": str(a.get("tracePath") or ""),
        "a_artifact_snapshot": str(a.get("artifactSnapshotUrl") or ""),
        "a_screencast": str((recordings.get("A") or {}).get("videoPath") or ""),
        "a_score_delivery": str(delivery_a.get("score") or ""),
        "a_desc_delivery": str(delivery_a.get("description") or "").strip(),
        "b_session_id": str(b.get("sessionId") or ""),
        "b_trace_file": str(b.get("tracePath") or ""),
        "b_artifact_snapshot": str(b.get("artifactSnapshotUrl") or ""),
        "b_screencast": str((recordings.get("B") or {}).get("videoPath") or ""),
        "b_score_delivery": str(delivery_b.get("score") or ""),
        "b_desc_delivery": str(delivery_b.get("description") or "").strip(),
        "gsb_verdict": str(draft.get("verdict") or ""),
        "gsb_reason": str(draft.get("reason") or "").strip(),
    }


def inspect_trace(path: Path, expected_prompt: str, expected_session: str) -> dict:
    result = {
        "path": str(path),
        "ok": False,
        "lineCount": 0,
        "sessionIds": [],
        "humanPromptCount": 0,
        "finalStopReason": "",
        "errors": [],
    }
    if not path.is_file():
        result["errors"].append("文件不存在")
        return result
    events: list[dict] = []
    try:
        with path.open(encoding="utf-8", errors="replace") as stream:
            for line_number, raw in enumerate(stream, 1):
                if not raw.strip():
                    continue
                try:
                    event = json.loads(raw)
                except json.JSONDecodeError as exc:
                    result["errors"].append(f"第 {line_number} 行 JSONL 无法解析: {exc}")
                    continue
                if isinstance(event, dict):
                    events.append(event)
    except OSError as exc:
        result["errors"].append(f"读取失败: {exc}")
        return result

    result["lineCount"] = len(events)
    expected_session = expected_session.strip()
    session_values = [
        str(event.get("sessionId") or event.get("session_id") or "").strip()
        for event in events
        if str(event.get("sessionId") or event.get("session_id") or "").strip()
    ]
    session_ids = set(session_values)
    result["expectedSessionId"] = expected_session
    result["sessionIdEventCount"] = len(session_values)
    result["sessionIds"] = sorted(session_ids)
    if not expected_session:
        result["errors"].append("状态或 Excel 缺少预期 SessionID")
    elif not session_values:
        result["errors"].append("轨迹文件内容不包含 SessionID")
    elif session_ids != {expected_session}:
        result["errors"].append(
            f"轨迹文件 SessionID 与预期不一致: {sorted(session_ids)} != {expected_session}"
        )

    prompt_events = []
    for index, event in enumerate(events):
        if event.get("type") != "user" or event.get("isMeta"):
            continue
        content = (event.get("message") or {}).get("content")
        if isinstance(content, str):
            prompt_events.append((index, content))
    result["humanPromptCount"] = len(prompt_events)
    if len(prompt_events) != 1:
        result["errors"].append(f"真人 Prompt 数量应为 1，实际 {len(prompt_events)}")
    elif prompt_events[0][1].strip() != expected_prompt.strip():
        result["errors"].append("轨迹 Prompt 与提示词文件不一致")

    final_assistant_index = -1
    for index, event in enumerate(events):
        if event.get("type") == "assistant":
            stop_reason = str((event.get("message") or {}).get("stop_reason") or "")
            if stop_reason:
                result["finalStopReason"] = stop_reason
                final_assistant_index = index
    if result["finalStopReason"] != "end_turn":
        result["errors"].append(f"最终 assistant stop_reason 不是 end_turn: {result['finalStopReason'] or '缺失'}")
    if final_assistant_index >= 0:
        for event in events[final_assistant_index + 1 :]:
            event_type = str(event.get("type") or "")
            subtype = str(event.get("subtype") or "")
            if event_type in {"assistant", "user", "error"} or subtype == "api_error":
                result["errors"].append(f"end_turn 后仍有异常事件: {event_type}/{subtype}")
                break
    for event in events:
        event_type = str(event.get("type") or "")
        subtype = str(event.get("subtype") or "")
        if event_type == "error" or subtype == "api_error":
            result["errors"].append(f"轨迹包含 API/网络错误事件: {event_type}/{subtype}")
            break
    result["ok"] = not result["errors"]
    return result


def inspect_media(path: Path, expected_kind: str) -> dict:
    result = {
        "path": str(path),
        "ok": False,
        "sizeBytes": None,
        "sha256": "",
        "extension": path.suffix.lower(),
        "width": None,
        "height": None,
        "durationSeconds": None,
        "errors": [],
    }
    if not path.is_file():
        result["errors"].append("文件不存在")
        return result
    result["sizeBytes"] = path.stat().st_size
    result["sha256"] = sha256_file(path)
    if expected_kind == "video":
        if path.suffix.lower() not in VIDEO_EXTS:
            result["errors"].append(f"视频格式不支持: {path.suffix}")
        if result["sizeBytes"] > VIDEO_MAX:
            result["errors"].append("视频超过 500 MB")
        if shutil.which("ffprobe"):
            probe = run(
                [
                    "ffprobe",
                    "-v",
                    "error",
                    "-select_streams",
                    "v:0",
                    "-show_entries",
                    "stream=width,height:format=duration",
                    "-of",
                    "json",
                    str(path),
                ],
                timeout=30,
            )
            if probe.returncode != 0:
                result["errors"].append(probe.stderr.strip() or "ffprobe 失败")
            else:
                try:
                    payload = json.loads(probe.stdout)
                    stream = (payload.get("streams") or [{}])[0]
                    fmt = payload.get("format") or {}
                    result["width"] = int(stream.get("width") or 0)
                    result["height"] = int(stream.get("height") or 0)
                    result["durationSeconds"] = float(fmt.get("duration") or 0)
                except Exception as exc:
                    result["errors"].append(f"ffprobe 输出无法解析: {exc}")
                if (result["width"], result["height"]) != (1280, 720):
                    result["errors"].append(f"视频分辨率不是 1280x720: {result['width']}x{result['height']}")
                if not result["durationSeconds"] or result["durationSeconds"] > 90.5:
                    result["errors"].append(f"视频时长无效或超过 90 秒: {result['durationSeconds']}")
        else:
            result["errors"].append("缺少 ffprobe")
    else:
        if path.suffix.lower() != ".jsonl":
            result["errors"].append(f"轨迹格式不是 .jsonl: {path.suffix}")
        if result["sizeBytes"] > TRACE_MAX:
            result["errors"].append("轨迹超过 27 MB")
    result["ok"] = not result["errors"]
    return result


def read_excel(path: Path, schema: dict) -> dict:
    result = {
        "path": str(path),
        "ok": False,
        "headers": [],
        "values": {},
        "missingRequired": [],
        "mismatches": [],
        "errors": [],
    }
    if load_workbook is None:
        result["errors"].append(f"缺少 openpyxl: {OPENPYXL_ERROR}")
        return result
    if not path.is_file():
        result["errors"].append("Excel 不存在")
        return result
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook["GSB提交"] if "GSB提交" in workbook.sheetnames else workbook.active
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        values_row = next(sheet.iter_rows(min_row=2, max_row=2), [])
        values = {}
        for index, header in enumerate(headers):
            values[header] = str(values_row[index].value or "").strip() if index < len(values_row) else ""
        result["headers"] = headers
        result["values"] = values
        workbook.close()
    except Exception as exc:
        result["errors"].append(f"Excel 读取失败: {exc}")
        return result

    expected_headers = [str(item.get("label") or "") for item in sorted(schema.get("fields") or [], key=lambda item: item.get("order", 0))]
    if headers != expected_headers:
        result["errors"].append("Excel 表头顺序与页面字段快照不一致")
    for field in schema.get("fields") or []:
        label = str(field.get("label") or "")
        if field.get("required") and not values.get(label, "").strip():
            result["missingRequired"].append(label)
    result["ok"] = not result["errors"] and not result["missingRequired"]
    return result


def compare_excel(excel: dict, expected: dict, schema: dict) -> list[str]:
    mismatches = []
    by_label = {str(item.get("label") or ""): item.get("key") for item in schema.get("fields") or []}
    for key, target in expected.items():
        label = next((label for label, candidate in by_label.items() if candidate == key), "")
        if not label:
            continue
        actual = str((excel.get("values") or {}).get(label) or "").strip()
        target = str(target or "").strip()
        if key in {"a_trace_file", "b_trace_file", "a_screencast", "b_screencast"}:
            actual_resolved = str(Path(actual).expanduser().resolve()) if actual else ""
            target_resolved = str(Path(target).expanduser().resolve()) if target else ""
            if actual_resolved != target_resolved:
                mismatches.append(f"{label}: Excel={actual or '<空>'} / state={target or '<空>'}")
        elif actual != target:
            mismatches.append(f"{label}: Excel 与 state/draft 不一致")
    return mismatches


def verify_remote(state: dict, task_root: Path) -> dict:
    result = {"ok": False, "heads": {}, "parents": {}, "errors": []}
    remote = str(state.get("remoteUrl") or "")
    repo_url = str(state.get("repoUrl") or "")
    if not remote:
        result["errors"].append("state 缺少 remoteUrl")
        return result
    probe = run(["git", "ls-remote", "--heads", remote], timeout=60, env=github_env(require_proxy=True))
    if probe.returncode != 0:
        result["errors"].append(probe.stderr.strip() or "git ls-remote 失败")
        return result
    heads = {}
    for line in probe.stdout.splitlines():
        parts = line.split()
        if len(parts) == 2 and parts[1].startswith("refs/heads/"):
            heads[parts[1].split("refs/heads/", 1)[1]] = parts[0]
    result["heads"] = heads
    main_branches = [name for name in ("main", "master") if name in heads]
    if len(main_branches) != 1 or set(heads) != set(main_branches + ["A", "B"]):
        result["errors"].append(f"远端分支不符合主分支+A/B: {sorted(heads)}")
        return result
    initial = str(state.get("initialSnapshot") or "").lower()
    if heads.get(main_branches[0]) != initial:
        result["errors"].append("远端主分支不是初始快照")
    for side in ("A", "B"):
        expected = str(((state.get("sides") or {}).get(side) or {}).get("artifactSnapshot") or "").lower()
        if heads.get(side) != expected:
            result["errors"].append(f"{side} 分支与 state.artifactSnapshot 不一致")

    owner, repo = parse_remote(repo_url)
    for sha in [initial] + [str(((state.get("sides") or {}).get(side) or {}).get("artifactSnapshot") or "") for side in ("A", "B")]:
        if not sha:
            continue
        local_repo = task_root / "source" / "origin"
        parent = ""
        if local_repo.is_dir():
            parent_probe = run(["git", "-C", str(local_repo), "show", "-s", "--format=%P", sha], timeout=20)
            if parent_probe.returncode == 0:
                parent = parent_probe.stdout.strip().split()[0] if parent_probe.stdout.strip() else ""
        if not parent and owner and repo and shutil.which("gh"):
            api_probe = run(
                ["gh", "api", f"repos/{owner}/{repo}/commits/{sha}", "--jq", ".parents[0].sha"],
                timeout=60,
                env=github_env(require_proxy=True),
            )
            if api_probe.returncode == 0:
                parent = api_probe.stdout.strip()
        if parent:
            result["parents"][sha] = parent
    for side in ("A", "B"):
        sha = str(((state.get("sides") or {}).get(side) or {}).get("artifactSnapshot") or "").lower()
        if result["parents"].get(sha) != initial:
            result["errors"].append(f"{side} 产物快照父提交不是初始快照")
    result["ok"] = not result["errors"]
    return result


def validate_reason_similarity_approval(task_root: Path, reason: str, dedup: dict) -> tuple[bool, dict]:
    if str(dedup.get("decision") or "") != "REVIEW_REQUIRED":
        return False, {"error": "人工审批只能覆盖 REVIEW_REQUIRED，不能覆盖 EXACT 或 SIMILAR"}
    path = task_root / "workspace" / "评审文件" / "pre-submit" / "reason-similarity-approval.json"
    if not path.is_file():
        return False, {"error": "缺少理由相似度人工审批文件", "path": str(path)}
    try:
        approval = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return False, {"error": f"理由相似度人工审批文件不可读: {exc}", "path": str(path)}
    reason_sha = hashlib.sha256(reason.strip().encode("utf-8")).hexdigest()
    matches = [item for item in dedup.get("matches") or [] if isinstance(item, dict) and item.get("decision") == "REVIEW_REQUIRED"]
    match_ids = [int(item.get("id")) for item in matches if str(item.get("id") or "").isdigit()]
    approved_ids = [int(value) for value in (approval.get("reviewedHistoryIds") or []) if str(value).isdigit()]
    ok = (
        approval.get("approvalKind") == "manual"
        and approval.get("scope") == "reason-similarity"
        and str(approval.get("approvedBy") or "").strip().casefold() == AUTO_APPROVER
        and approval.get("reviewDecision") == "REVIEW_REQUIRED"
        and str(approval.get("reasonSha256") or "") == reason_sha
        and sorted(set(match_ids)) == sorted(set(approved_ids))
    )
    evidence = {"path": str(path), "approvedBy": approval.get("approvedBy"), "reasonSha256": approval.get("reasonSha256"), "reviewedHistoryIds": approved_ids, "currentMatches": match_ids}
    if not ok:
        evidence["error"] = "人工审批字段、理由哈希或命中历史列表不匹配"
    return ok, evidence


def assess_reason_quality(draft: dict, evidence_doc: dict | None = None) -> dict:
    """Validate the observed G5 requirements before a remote submission."""
    reason = str(draft.get("reason") or "").strip()
    claims = draft.get("claims") if isinstance(draft.get("claims"), list) else []
    negative = [item for item in claims if isinstance(item, dict) and item.get("polarity") == "negative"]
    errors: list[str] = []
    warnings: list[str] = []
    markdown_errors: list[str] = []
    language_errors: list[str] = []
    excluded_evidence_errors: list[str] = []
    def normalize(value) -> str:
        return re.sub(r"\s+", "", str(value or "")).strip()

    clean_reason = normalize(reason)
    layer_coverage: dict[str, bool] = {}
    if not negative:
        errors.append("draft 缺少 polarity=negative 的 claim，无法审核客观后果")
    for side in ("A", "B"):
        side_claims = [item for item in claims if isinstance(item, dict) and item.get("side") == side]
        if not any(item.get("polarity") == "positive" for item in side_claims):
            errors.append(f"{side} 侧缺少正面 claim")
        if not any(item.get("polarity") == "negative" for item in side_claims):
            errors.append(f"{side} 侧缺少负面 claim")
        for claim_type in ("process", "artifact"):
            key = f"{side}-{claim_type}"
            candidates = [item for item in side_claims if item.get("type") == claim_type]
            selected = False
            for claim in candidates:
                claim_text = str(claim.get("text") or "").strip()
                if normalize(claim_text) not in clean_reason:
                    continue
                if claim_type == "process":
                    if PROCESS_ACTION_RE.search(claim_text) and PROCESS_LOCATOR_RE.search(claim_text):
                        selected = True
                        break
                elif ARTIFACT_OUTCOME_RE.search(claim_text):
                    selected = True
                    break
            layer_coverage[key] = selected
            if not selected:
                if claim_type == "process":
                    errors.append(
                        f"{side} 侧缺少写入 GSB 理由的具体过程事实；需同时包含读取/修改/执行等实际动作，"
                        "以及文件、命令、步骤或需求等定位点"
                    )
                else:
                    errors.append(
                        f"{side} 侧缺少写入 GSB 理由的具体产物事实；需包含返回、缺少、未实现、失败、"
                        "写入、生成或通过等可观察结果"
                    )
    for index, claim in enumerate(negative, 1):
        trigger = str(claim.get("trigger") or "").strip()
        consequence = str(claim.get("objectiveConsequence") or "").strip()
        if not trigger:
            errors.append(f"负面 claim {index} 缺少 trigger")
        elif normalize(trigger) not in clean_reason:
            errors.append(f"负面 claim {index} 的 trigger 未写进 GSB 理由: {trigger}")
        if not consequence:
            errors.append(f"负面 claim {index} 缺少 objectiveConsequence 客观后果")
            continue
        if not OBJECTIVE_CONSEQUENCE_RE.search(consequence):
            errors.append(f"负面 claim {index} 的客观后果不可见: {consequence}")
        if normalize(consequence) not in clean_reason:
            errors.append(f"负面 claim {index} 的 objectiveConsequence 未写进 GSB 理由: {consequence}")
    if negative and not OBJECTIVE_EVIDENCE_RE.search(reason):
        errors.append("负面理由缺少文件名、函数名、命令、报错、退出码或接口状态等可核对证据")
    field_factor = FIELD_FACTOR_RE.search(reason)
    if field_factor:
        errors.append(
            "GSB 理由不得引用录屏、视频、截图、浏览器、测试设备、运行环境、验收宿主或采集现场因素；没有例外: "
            f"{field_factor.group(0)}"
        )
    acceptance = NON_CONTAINER_TEST_ARTIFACT_RE.search(reason)
    if acceptance:
        errors.append(
            "GSB 理由不得引用非容器内生成的测试文件、测试脚本、自测脚本或验收脚本；"
            "这些本地（容器外）内容不参与 GSB 文案，没有例外: "
            f"{acceptance.group(0)}"
        )
    banned_reason = BANNED_REASON_PATTERN.search(reason)
    if banned_reason:
        errors.append(
            "GSB 理由禁止使用“落在……”式收束句式；请直接写“因此选择 B 侧方案”或“B 侧方案更好”: "
            f"{banned_reason.group(0)}"
        )
    style_helpers = _load_reason_style_helpers()
    if style_helpers is not None:
        errors.extend(style_helpers.reason_style_errors(reason))
        errors.extend(style_helpers.reason_flow_errors(reason))
        attribution = style_helpers.SOURCE_ATTRIBUTION_RE.search(reason)
        if attribution:
            errors.append(
                "GSB 理由不要交代信息来源，直接写做了什么、看到什么，例如“请求了登录接口，返回404”: "
                f"{attribution.group(0)}"
            )
        warnings.extend(style_helpers.reason_style_warnings(reason))
        markdown_errors = style_helpers._validate_reason_markdown(reason)
        language_errors = style_helpers.reason_language_errors(reason)
        if evidence_doc is not None:
            excluded_evidence_errors = style_helpers.evaluation_excluded_reason_errors(
                reason, draft, evidence_doc
            )
        errors.extend(markdown_errors)
        errors.extend(language_errors)
        errors.extend(excluded_evidence_errors)
    else:
        errors.append("GSB 文案校验器 gsb_tools 不可用，拒绝提交")
        for word, better in {"落库": "入库", "闭环": "完整覆盖", "根因": "原因"}.items():
            if word in reason:
                errors.append(f"GSB 理由禁用“{word}”，请改写成“{better}”: {word}")
        masked = reason
        for label in ("A 侧方案", "B 侧方案"):
            masked = masked.replace(label, label.replace(" ", ""))
        spacing = re.search(r"[A-Za-z0-9]\s|\s[A-Za-z0-9]", masked)
        if spacing:
            errors.append(
                "GSB 理由的数字与英文两侧不要留空格，直接连写，例如“计数从2变22”: "
                f"{spacing.group(0)!r}"
            )
        if re.search(r"(?:录屏|视频|截图|浏览器|测试设备|运行环境|验收宿主)", reason):
            errors.append("GSB 理由不得引用录屏、视频、截图、浏览器、测试设备、运行环境或验收宿主")
        if "`" in reason or re.search(r"(?m)^\s*(?:#{1,6}\s|[-+*]\s|\d+[.)]\s|>\s?\|.*\|\s*$)", reason):
            errors.append("GSB 理由不得包含 Markdown 语法")
        if len(reason) > 240 or any(len(part) > 56 for part in re.split(r"[。；！？!?]", reason)):
            errors.append("GSB 理由单句过长或句子不通顺")
    if MACHINE_STEP_RE.search(reason):
        warnings.append("GSB 理由出现三位以上的机器式步骤序号，建议改成自然触发节点")
    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "negativeClaims": len(negative),
        "objectiveEvidence": bool(OBJECTIVE_EVIDENCE_RE.search(reason)),
        "fieldFactor": field_factor.group(0) if field_factor else "",
        "markdownErrors": markdown_errors,
        "languageErrors": language_errors,
        "excludedEvidenceErrors": excluded_evidence_errors,
        "layerCoverage": {"ok": all(layer_coverage.values()), "sides": layer_coverage},
        "reason": reason,
    }


def classify_change_volume_line_gate(
    checks: list[dict],
    code_change: dict,
    *,
    skip_remote: bool,
) -> dict:
    """Identify whether the line gate is the only policy-approved blocker."""
    line_gate_check_ids = {"code-volume-hard-a", "code-volume-hard-b"}
    failed_line_checks = [
        item for item in checks
        if item.get("id") in line_gate_check_ids and not item.get("ok")
    ]
    failed_sides = [
        side for side in ("A", "B")
        if not bool((((code_change.get("sides") or {}).get(side) or {}).get("hardOk")))
    ]
    other_failed_blockers = [
        item for item in checks
        if item.get("severity") == "blocker" and not item.get("ok")
        and item.get("id") not in line_gate_check_ids
    ]
    review_available = bool(
        not skip_remote
        and code_change
        and not code_change.get("errors")
        and code_change.get("okHygiene")
        and (code_change.get("sides") or {})
    )
    only_blocker = bool(
        failed_sides
        and failed_line_checks
        and review_available
        and not other_failed_blockers
    )
    return {
        "failedSides": failed_sides,
        "onlyBlocker": only_blocker,
        "failedCheckIds": [str(item.get("id") or "") for item in failed_line_checks],
        "otherBlockingCheckIds": [str(item.get("id") or "") for item in other_failed_blockers],
        "reviewAvailable": review_available,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Read-only GSB submission preflight")
    parser.add_argument("--task-root", type=Path)
    parser.add_argument("--excel", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--skip-remote", action="store_true")
    args = parser.parse_args()

    task_root = find_task_root(args.task_root)
    state_path = task_root / "monitor" / "state.json"
    state = load_json(state_path)
    draft_path = task_root / "monitor" / "gsb-draft.json"
    if not draft_path.is_file():
        # 0917 技能把草稿写在任务根目录；独立提交技能同时兼容新旧布局。
        draft_path = task_root / "gsb-draft.json"
    draft = load_json(draft_path)
    excel_path = (args.excel or task_root / "workspace" / "评审文件" / "交付表.xlsx").expanduser().resolve()
    output_dir = (args.output_dir or task_root / "workspace" / "评审文件" / "pre-submit").expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    schema = page_schema()
    prompt_path = Path(str(state.get("promptPath") or "")).expanduser()
    prompt_text = prompt_path.read_text(encoding="utf-8") if prompt_path.is_file() else ""
    expected = derived_expected({**state, "promptText": prompt_text}, draft)

    current_sessions = {
        str(((state.get("sides") or {}).get(side) or {}).get("sessionId") or "")
        for side in ("A", "B")
    } - {""}
    current_result = load_json(task_root / "workspace" / "评审文件" / "pre-submit" / "submission-result.json")
    exclude_ids = {int(current_result.get("submissionNo"))} if str(current_result.get("submissionNo") or "").isdigit() else set()
    prompt_history = {}
    prompt_dedup = {}
    reason_history = {}
    reason_dedup = {}
    history_error = ""
    history_path = task_root / "monitor" / "gsb-prompt-history.json"
    dedup_path = task_root / "monitor" / "prompt-dedup-review.json"
    reason_history_path = task_root / "monitor" / "gsb-reason-history.json"
    reason_dedup_path = task_root / "monitor" / "gsb-reason-dedup-review.json"
    reason_text = str(draft.get("reason") or expected.get("gsb_reason") or "").strip()
    try:
        # Submission preflight forces a live refresh, then persists the complete
        # local history cache for future dry-runs and offline review.
        prompt_history = fetch_gsb_prompt_history(
            current_sessions,
            exclude_ids,
            force_refresh=True,
        )
        history_error = str(prompt_history.get("fetchError") or "")
        prompt_dedup = assess_prompt_dedup(prompt_text, prompt_history)
        reason_history = {
            **prompt_history,
            "items": [
                {key: item.get(key) for key in (
                    "id", "status", "submittedAt", "gsbReason", "questionType", "difficulty",
                    "languages", "repoId", "aSessionId", "bSessionId",
                )}
                for item in prompt_history.get("items") or []
                if str(item.get("gsbReason") or "").strip()
            ],
        }
        reason_history["total"] = len(reason_history["items"])
        reason_dedup = assess_gsb_reason_dedup(reason_text, reason_history)
        for path, value in (
            (history_path, prompt_history),
            (dedup_path, prompt_dedup),
            (reason_history_path, reason_history),
            (reason_dedup_path, reason_dedup),
        ):
            path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except Exception as exc:
        history_error = str(exc)

    checks: list[dict] = []
    blockers: list[str] = []
    warnings: list[str] = []

    def add(check_id: str, ok: bool, message: str, *, severity: str = "blocker", evidence: dict | None = None) -> None:
        item = {"id": check_id, "ok": bool(ok), "severity": severity, "message": message, "evidence": evidence or {}}
        checks.append(item)
        if not ok:
            (blockers if severity == "blocker" else warnings).append(message)

    add("state", state.get("status") in {"complete", "recorded", "gsb_ready"}, f"任务状态: {state.get('status') or '缺失'}", evidence={"path": str(state_path)})
    add("prompt", prompt_path.is_file() and bool(prompt_text.strip()), "Prompt 文件存在且非空", evidence={"path": str(prompt_path)})
    if prompt_path.is_file() and state.get("promptSha256"):
        add("prompt-sha", sha256_file(prompt_path) == str(state.get("promptSha256")), "Prompt SHA-256 与 state 一致")
    add("gsb-history-live", bool(prompt_history) and not history_error, "提交前已刷新历史 GSB 文案缓存", evidence={"error": history_error, "cachePath": prompt_history.get("cachePath", ""), "cacheFresh": prompt_history.get("cacheFresh", False), "cacheHit": prompt_history.get("cacheHit", False), "total": prompt_history.get("total", 0)})
    unresolved_history = unresolved_history_ids(prompt_history)
    add(
        "gsb-history-unresolved",
        not unresolved_history,
        "平台反馈的历史记录已具备理由文本，可执行 B-5 去重",
        evidence={"ids": unresolved_history, "manualCachePath": prompt_history.get("manualCachePath", str(gsb_history_manual_cache_path()))},
    )
    add("prompt-history", bool(prompt_history), "历史 GSB 提示词列表已抽取", evidence={"path": str(history_path), "total": prompt_history.get("total", 0), "skippedCurrentSubmissions": prompt_history.get("skippedCurrentSubmissions", [])})
    add("prompt-dedup", bool(prompt_dedup) and prompt_dedup.get("decision") == "UNIQUE", f"历史 GSB 提示词去重: {prompt_dedup.get('decision') or 'BLOCKED'}", evidence={"path": str(dedup_path), "decision": prompt_dedup.get("decision"), "matches": prompt_dedup.get("matches", []), "error": history_error})
    add("gsb-reason-history", bool(reason_history), "历史 GSB 理由列表已抽取", evidence={"path": str(reason_history_path), "total": reason_history.get("total", 0)})
    reason_decision = str(reason_dedup.get("decision") or "BLOCKED")
    reason_message = f"历史 GSB 理由去重: {reason_decision}"
    if reason_dedup.get("rewriteInstruction"):
        reason_message += f"；{reason_dedup.get('rewriteInstruction')}"
    manual_reason_ok = False
    manual_reason_evidence = {}
    if reason_decision == "REVIEW_REQUIRED":
        manual_reason_ok, manual_reason_evidence = validate_reason_similarity_approval(task_root, reason_text, reason_dedup)
        if manual_reason_ok:
            reason_message += "；已由配置里的审批人人工审核通过"
    add(
        "gsb-reason-dedup",
        reason_decision == "UNIQUE" or manual_reason_ok,
        reason_message,
        evidence={
            "path": str(reason_dedup_path),
            "decision": reason_decision,
            "matches": reason_dedup.get("matches", []),
            "rewriteInstruction": reason_dedup.get("rewriteInstruction", ""),
            "manualApproval": manual_reason_evidence,
        },
    )
    add("draft", bool(draft), "gsb-draft.json 存在", evidence={"path": str(draft_path)})
    evidence_doc = load_json(task_root / "monitor" / "evidence.json")
    reason_quality = assess_reason_quality(draft, evidence_doc)
    add(
        "reason-quality",
        reason_quality["ok"],
        "GSB 理由分别覆盖 A/B 过程与产物，包含触发节点、客观后果和可核对证据",
        evidence=reason_quality,
    )
    for warning in reason_quality.get("warnings") or []:
        add("reason-ai-style", False, warning, severity="warning", evidence=reason_quality)

    style_helpers = _load_reason_style_helpers()
    if style_helpers is not None:
        delivery_quality = style_helpers.validate_delivery(draft, evidence_doc, reason_text)
    else:
        delivery_quality = {"ok": False, "errors": ["GSB 文案校验器 gsb_tools 不可用，无法审核交付完整性"], "warnings": []}
    add(
        "delivery-quality",
        delivery_quality["ok"],
        "A/B 交付完整性打分为 1~5 整数，描述只谈完整性、按实际情况独立撰写且未照抄 GSB 理由",
        evidence=delivery_quality,
    )
    for warning in delivery_quality.get("warnings") or []:
        add("delivery-verdict-consistency", False, warning, severity="warning", evidence=delivery_quality)
    delivery_texts = {
        side: str(((draft.get("delivery") or {}).get(side) or {}).get("description") or "").strip()
        for side in ("A", "B")
    } if isinstance(draft.get("delivery"), dict) else {"A": "", "B": ""}
    delivery_dedup = assess_delivery_dedup(delivery_texts, prompt_history) if prompt_history else {"decision": "BLOCKED"}
    delivery_dedup_path = task_root / "monitor" / "gsb-delivery-dedup-review.json"
    delivery_dedup_path.write_text(json.dumps(delivery_dedup, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    delivery_decision = str(delivery_dedup.get("decision") or "BLOCKED")
    add(
        "delivery-dedup",
        delivery_decision in {"UNIQUE", "REVIEW_REQUIRED"},
        f"G12 交付完整性描述历史去重: {delivery_decision}",
        evidence={"path": str(delivery_dedup_path), "decision": delivery_decision},
    )
    if delivery_decision == "REVIEW_REQUIRED":
        add("delivery-dedup-review", False, "交付完整性描述与历史描述有 8 字以上公共片段，建议换一种说法", severity="warning", evidence={"path": str(delivery_dedup_path)})

    excel = read_excel(excel_path, schema)
    add("excel", excel.get("ok"), "Excel 字段存在且必填项完整", evidence=excel)
    if excel.get("missingRequired"):
        add("excel-required", False, f"Excel 缺少必填字段: {', '.join(excel['missingRequired'])}")
    mismatches = compare_excel(excel, expected, schema) if excel.get("values") else []
    add("excel-values", not mismatches, "Excel 与 task root 当前值一致", evidence={"mismatches": mismatches})
    if mismatches:
        blockers.extend(mismatches)
    repro_field = next((field for field in schema.get("fields") or [] if field.get("key") == "repro_level"), {})
    repro_options = {str(item) for item in repro_field.get("options") or []}
    repro_actual = str((excel.get("values") or {}).get("环境可复现等级") or "").strip()
    repro_ok = bool(repro_options) and repro_actual in repro_options
    add("excel-repro-enum", repro_ok, "环境可复现等级使用页面允许值", evidence={"actual": repro_actual, "options": sorted(repro_options)})

    trace_results = {}
    for side in ("A", "B"):
        side_state = (state.get("sides") or {}).get(side) or {}
        trace = Path(str(side_state.get("tracePath") or "")).expanduser()
        info = inspect_trace(trace, prompt_text, str(side_state.get("sessionId") or ""))
        info["upload"] = inspect_media(trace, "trace")
        trace_results[side] = info
        add(f"trace-{side.lower()}", info.get("ok"), f"{side} 轨迹结构校验", evidence=info)
        add(f"trace-upload-{side.lower()}", info["upload"].get("ok"), f"{side} 轨迹上传规格", evidence=info["upload"])

    video_results = {}
    for side in ("A", "B"):
        video_path = Path(str(((state.get("recordings") or {}).get(side) or {}).get("videoPath") or "")).expanduser()
        info = inspect_media(video_path, "video")
        browser_result = task_root / "monitor" / "recording" / side.lower() / "web-otty" / "browser-result.json"
        browser_status = str(load_json(browser_result).get("status") or "")
        info["browserResult"] = str(browser_result) if browser_result.is_file() else ""
        info["browserStatus"] = browser_status
        recording = ((state.get("recordings") or {}).get(side) or {})
        expected_failure = bool(recording.get("expectedAppFailure"))
        observed_failure = bool(recording.get("observedAppFailure"))
        recording_ok = bool(recording.get("ok"))
        outcome_consistent = expected_failure == observed_failure
        app_outcome_consistent = str(recording.get("appOutcome") or "") == ("failure" if observed_failure else "success")
        info["recordingOk"] = recording_ok
        info["expectedAppFailure"] = expected_failure
        info["observedAppFailure"] = observed_failure
        info["outcomeConsistent"] = outcome_consistent
        if not recording_ok:
            info["errors"].append("state.recordings.ok 不是 true")
            info["ok"] = False
        if not outcome_consistent:
            info["errors"].append("录制失败策略与真实退出结果不一致")
            info["ok"] = False
        if not app_outcome_consistent:
            info["errors"].append(f"appOutcome 与 observedAppFailure 不一致: {recording.get('appOutcome')}")
            info["ok"] = False
        if browser_result.is_file() and browser_status != "passed" and not (expected_failure and observed_failure):
            info["errors"].append(f"browser-result 状态不是 passed: {browser_status}")
            info["ok"] = False
        video_results[side] = info
        add(f"recording-{side.lower()}", recording_ok and outcome_consistent and app_outcome_consistent, f"{side} 录制结果与失败策略一致", evidence=info)
        add(f"video-{side.lower()}", info.get("ok"), f"{side} 视频上传规格与真实运行结果", evidence=info)

    code_change = {}
    if not args.skip_remote:
        remote = verify_remote(state, task_root)
        add("remote-git", remote.get("ok"), "远端 Git 分支和 commit 父关系", evidence=remote)
        if remote.get("ok"):
            code_change = verify_code_change(state, remote)
            add(
                "repo-hygiene",
                code_change.get("okHygiene", False),
                "A/B 产物未包含依赖目录、构建产物或锁文件",
                evidence={"badPaths": code_change.get("badPaths", [])[:100]},
            )
            for side in ("A", "B"):
                side_info = (code_change.get("sides") or {}).get(side) or {}
                add(
                    f"code-volume-hard-{side.lower()}",
                    side_info.get("hardOk", False),
                    f"{side} 非测试业务代码改动至少 {CODE_VOLUME_HARD_MIN} 行",
                    evidence=side_info,
                )
                add(
                    f"code-volume-target-{side.lower()}",
                    side_info.get("targetOk", False),
                    f"{side} 达到建议改动量：至少 {CODE_VOLUME_TARGET} 行且 3 个业务文件",
                    severity="warning",
                    evidence=side_info,
                )
        else:
            add("repo-hygiene", False, "远端 Git 不可用，无法检查依赖目录和锁文件", evidence=remote)
            add("code-volume", False, "远端 Git 不可用，无法检查 G11 改动量", evidence=remote)

    required = [field for field in schema.get("fields") or [] if field.get("required")]
    form = schema.get("form") or {}
    add(
        "schema-page",
        len(schema.get("fields") or []) == int(form.get("fieldCount") or 0)
        and len(required) == int(form.get("requiredCount") or 0),
        f"页面字段快照为 {form.get('fieldCount')} 字段/{form.get('requiredCount')} 必填",
    )
    reason = str(draft.get("reason") or "")
    reason_nonwhite = len(re.sub(r"\s+", "", reason))
    state_versions = {side: str(((state.get("sides") or {}).get(side) or {}).get("harnessVersion") or "") for side in ("A", "B")}
    add(
        "harness-version",
        all(version.startswith(HARNESS_VERSION) for version in state_versions.values()),
        f"Harness 版本固定为 {HARNESS_VERSION}",
        evidence=state_versions,
    )
    add("reason-page-min", reason_nonwhite >= 60, "GSB 理由满足页面最少 60 字", severity="blocker", evidence={"nonWhiteLength": reason_nonwhite})
    add("reason-local-range", 150 <= reason_nonwhite <= 240, "GSB 理由满足 0917 的 150–240 字", severity="warning", evidence={"nonWhiteLength": reason_nonwhite})
    add("reason-sides", bool(re.search(r"A|Ａ", reason)) and bool(re.search(r"B|Ｂ", reason)), "GSB 理由明确覆盖 A/B", severity="warning")
    test_count_match = re.search(
        r"(?:(?:后端|前端|API|接口|测试|断言|用例)[^。；，]{0,12}?\d+\s*(?:个|项|条|步|轮)|"
        r"\d+\s*(?:个|项|条|步|轮)[^。；，]{0,12}?(?:后端|前端|API|接口|测试|断言|用例))",
        reason,
        re.I,
    )
    add(
        "reason-no-test-counts",
        not bool(test_count_match),
        "GSB 理由用业务覆盖描述，不用测试或断言数量堆砌",
        severity="warning",
        evidence={"match": test_count_match.group(0) if test_count_match else ""},
    )
    layer_coverage = reason_quality.get("layerCoverage") or {}
    add(
        "reason-process-product",
        bool(layer_coverage.get("ok")),
        "GSB 理由分别覆盖 A/B 的过程与产物，且每层都有可核对事实",
        evidence=layer_coverage,
    )
    markdown_errors = reason_quality.get("markdownErrors") or []
    add(
        "reason-markdown",
        not markdown_errors,
        "GSB 理由不含标题、列表、代码块、行内代码、链接或强调标记",
        evidence={"errors": markdown_errors},
    )
    language_errors = reason_quality.get("languageErrors") or []
    add(
        "reason-high-school-language",
        not language_errors,
        "GSB 理由句子通顺，单句不超长，标点和表达达到高中语文水平",
        evidence={"errors": language_errors},
    )
    excluded_evidence_errors = reason_quality.get("excludedEvidenceErrors") or []
    add(
        "reason-evaluation-excluded",
        not excluded_evidence_errors,
        "GSB 理由没有引用 evaluationExcluded 的环境或工具噪声证据",
        evidence={"errors": excluded_evidence_errors},
    )

    line_gate_assessment = classify_change_volume_line_gate(
        checks,
        code_change,
        skip_remote=bool(args.skip_remote),
    )
    line_gate_failed_sides = line_gate_assessment["failedSides"]
    only_line_gate_blocker = line_gate_assessment["onlyBlocker"]
    review_payload = {
        "schemaVersion": 1,
        "id": CHANGE_VOLUME_APPROVAL_SCOPE,
        "initialSnapshot": str(state.get("initialSnapshot") or ""),
        "hardMinimumLines": CODE_VOLUME_HARD_MIN,
        "targetLines": CODE_VOLUME_TARGET,
        "failedSides": line_gate_failed_sides,
        "onlyBlocker": only_line_gate_blocker,
        "codeChange": code_change,
    }
    line_gate_review = {
        **review_payload,
        "reviewSha256": sha256_json(review_payload),
        "checkedAt": utc_now(),
    }
    line_gate_review_path = output_dir / "change-volume-review.json"
    line_gate_review_path.write_text(
        json.dumps(line_gate_review, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    uploads = {
        "a_trace": trace_results.get("A", {}).get("upload", {}),
        "a_video": video_results.get("A", {}),
        "b_trace": trace_results.get("B", {}).get("upload", {}),
        "b_video": video_results.get("B", {}),
    }
    schema_fields = {field.get("key"): field for field in schema.get("fields") or []}
    payload_fields = {
        str(schema_fields.get(key, {}).get("label") or key): value
        for key, value in expected.items()
        if str(schema_fields.get(key, {}).get("type") or "") not in {"attachment", "video"}
    }
    payload_uploads = {
        "A-轨迹文件": {"path": trace_results.get("A", {}).get("upload", {}).get("path", ""), "kind": "trace", "sha256": trace_results.get("A", {}).get("upload", {}).get("sha256", "")},
        "A-运行录屏": {"path": video_results.get("A", {}).get("path", ""), "kind": "video", "sha256": video_results.get("A", {}).get("sha256", "")},
        "B-轨迹文件": {"path": trace_results.get("B", {}).get("upload", {}).get("path", ""), "kind": "trace", "sha256": trace_results.get("B", {}).get("upload", {}).get("sha256", "")},
        "B-运行录屏": {"path": video_results.get("B", {}).get("path", ""), "kind": "video", "sha256": video_results.get("B", {}).get("sha256", "")},
    }
    submission_payload = {
        "schemaVersion": 1,
        "targetUrl": gsb_server() + "/app/gsb/submit",
        "taskRoot": str(task_root),
        "harnessVersion": HARNESS_VERSION,
        "deliverySheet": {
            "path": str(excel_path),
            "sha256": sha256_file(excel_path) if excel_path.is_file() else "",
        },
        "approvalRequired": only_line_gate_blocker,
        "approvalPolicy": (
            CHANGE_VOLUME_APPROVAL_SCOPE if only_line_gate_blocker
            else ("automatic" if not blockers else "blocked")
        ),
        "changeVolumeLineGate": {
            "required": only_line_gate_blocker,
            "failedSides": line_gate_failed_sides,
            "hardMinimumLines": CODE_VOLUME_HARD_MIN,
            "reviewPath": str(line_gate_review_path),
            "reviewSha256": line_gate_review["reviewSha256"],
        },
        "fields": payload_fields,
        "uploads": payload_uploads,
        "defaults": {"upload": False, "submit": False},
    }
    payload_path = output_dir / "submission-payload.json"
    payload_path.write_text(json.dumps(submission_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report = {
        "schemaVersion": 1,
        "checkedAt": utc_now(),
        "taskRoot": str(task_root),
        "page": {
            "url": gsb_server() + "/app/gsb/submit",
            "fieldCount": len(schema.get("fields") or []),
            "requiredCount": len(required),
            "schemaInspectedAt": schema.get("inspectedAt"),
        },
        "ok": not blockers,
        "status": (
            "pass" if not blockers
            else ("line_gate_approval_required" if only_line_gate_blocker else "blocked")
        ),
        "approvalRequirement": (
            CHANGE_VOLUME_APPROVAL_SCOPE if only_line_gate_blocker
            else ("none" if not blockers else "blocked")
        ),
        "lineGate": {
            "onlyBlocker": only_line_gate_blocker,
            "failedSides": line_gate_failed_sides,
            "reviewPath": str(line_gate_review_path),
            "reviewSha256": line_gate_review["reviewSha256"],
        },
        "blockers": blockers,
        "warnings": warnings,
        "checks": checks,
        "expected": expected,
        "excel": excel,
        "traces": trace_results,
        "videos": video_results,
        "uploads": uploads,
        "uploadPolicy": schema.get("uploadRules"),
        "promptHistoryPath": str(history_path),
        "promptDedupPath": str(dedup_path),
        "reasonHistoryPath": str(reason_history_path),
        "reasonDedupPath": str(reason_dedup_path),
        "historyCachePath": str(prompt_history.get("cachePath") or gsb_history_cache_path()),
        "historyFresh": bool(prompt_history.get("cacheFresh")) and not history_error,
        "unresolvedHistoryIds": unresolved_history_ids(prompt_history),
        "historyFetchError": history_error,
        "promptHistory": prompt_history,
        "promptDedup": prompt_dedup,
        "reasonHistory": reason_history,
        "reasonDedup": reason_dedup,
        "deliveryDedupPath": str(delivery_dedup_path),
        "deliveryDedup": delivery_dedup,
        "deliveryQuality": delivery_quality,
        "codeChange": code_change,
        "changeVolumeLineGate": line_gate_review,
        "submissionPayloadPath": str(payload_path),
        "submissionPayloadSha256": sha256_file(payload_path),
        "submissionPayload": submission_payload,
    }
    json_path = output_dir / "pre-submit-review.json"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    lines = [
        "# GSB 提交前审核",
        "",
        f"- 任务根目录：`{task_root}`",
        f"- 结论：{'通过，可进入上传授权阶段' if report['ok'] else '阻断，不能上传'}",
        f"- 页面字段：{report['page']['fieldCount']} 个，必填 {report['page']['requiredCount']} 个",
        f"- 审核时间：{report['checkedAt']}",
        "",
        "## 四个上传文件",
        "",
        "| 槽位 | 文件 | 大小 | SHA-256 | 结论 |",
        "|---|---|---:|---|---|",
    ]
    for key, info in uploads.items():
        size = info.get("sizeBytes")
        size_text = f"{size / 1024 / 1024:.2f} MB" if isinstance(size, int) else "-"
        lines.append(f"| `{key}` | `{info.get('path') or ''}` | {size_text} | `{info.get('sha256') or ''}` | {'通过' if info.get('ok') else '失败'} |")
    lines += ["", "## G11 代码改动量", ""]
    if code_change.get("sides"):
        lines.extend([
            f"- 硬底线：每个 A/B 至少 {code_change.get('hardMinimumLines', 10)} 行非测试业务代码。",
            f"- 建议目标：至少 {code_change.get('targetLines', 30)} 行且跨 {code_change.get('targetFiles', 3)} 个业务文件。",
        ])
        for side in ("A", "B"):
            side_info = (code_change.get("sides") or {}).get(side) or {}
            lines.append(
                f"- {side}：{side_info.get('lines', 0)} 行，{side_info.get('businessFileCount', 0)} 个业务文件，"
                f"硬门禁{'通过' if side_info.get('hardOk') else '阻断'}，建议目标{'通过' if side_info.get('targetOk') else '未达到'}。"
            )
        if code_change.get("badPaths"):
            lines.append(f"- 检测到依赖/构建/锁文件路径：`{len(code_change.get('badPaths') or [])}` 个，已阻断。")
    else:
        lines.append("- 未执行远端代码改动量检查。")
    lines += ["", "## 阻断项", ""]
    lines.extend([f"- {item}" for item in blockers] or ["- 无"])
    lines += ["", "## 历史 GSB 提示词去重", "", f"- 结论：`{prompt_dedup.get('decision') or 'BLOCKED'}`", f"- 历史提示词数：{prompt_history.get('total', 0)}", f"- 当前提交排除：{prompt_history.get('skippedCurrentSubmissions', [])}", f"- 本地缓存：`{prompt_history.get('cachePath') or gsb_history_cache_path()}`"]
    if prompt_dedup.get("matches"):
        for match in prompt_dedup["matches"]:
            lines.append(f"- `#{match.get('id')}` {match.get('decision')}，相似度 {match.get('similarity')}，最长连续片段 {match.get('longestCommonSubstringLength')} 字")
    else:
        lines.append("- 未发现精确、连续片段或高相似度重复。")
    lines += ["", "## 历史 GSB 理由去重", "", f"- 结论：`{reason_dedup.get('decision') or 'BLOCKED'}`", f"- 历史理由数：{reason_history.get('total', 0)}"]
    if reason_dedup.get("matches"):
        for match in reason_dedup["matches"]:
            lines.append(
                f"- `#{match.get('id')}` {match.get('decision')}，相似度 {match.get('similarityPercent')}%，"
                f"最长连续片段 {match.get('longestCommonSubstringLength')} 字，规则 `{match.get('rule')}`"
            )
        if reason_dedup.get("rewriteInstruction"):
            lines.append(f"- 重写要求：{reason_dedup.get('rewriteInstruction')}")
    else:
        lines.append("- 未发现 B-5 公共长片段、模板化片段或高相似度理由。")
    lines += ["", "## 警告", ""]
    lines.extend([f"- {item}" for item in warnings] or ["- 无"])
    lines += ["", "## 明细", ""]
    for check in checks:
        lines.append(f"- [{'PASS' if check['ok'] else 'FAIL'}] `{check['id']}`：{check['message']}")
    md_path = output_dir / "pre-submit-review.md"
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": report["status"],
        "ok": report["ok"],
        "json": str(json_path),
        "markdown": str(md_path),
        "blockers": blockers,
        "warnings": warnings,
        "approvalRequirement": report["approvalRequirement"],
        "lineGate": report["lineGate"],
        "submissionPayloadPath": report["submissionPayloadPath"],
        "submissionPayloadSha256": report["submissionPayloadSha256"],
        "historyCachePath": str(prompt_history.get("cachePath") or gsb_history_cache_path()),
        "historyFresh": bool(prompt_history.get("cacheFresh")) and not history_error,
        "reasonDedup": {
            "decision": reason_dedup.get("decision"),
            "matches": reason_dedup.get("matches", []),
            "rewriteInstruction": reason_dedup.get("rewriteInstruction", ""),
        },
    }, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
